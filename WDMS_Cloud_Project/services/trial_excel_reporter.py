"""Publish the three-column inventory workbook (《x年x月x日物流库存》) after readback validation."""

from __future__ import annotations

import os
import sys
import uuid
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence, Union

from openpyxl import Workbook, load_workbook

from core.numbers import (
    parse_nonnegative_integer,
    parse_nonnegative_quantity,
)
from errors import WDMSError
from services.workbook_table import InventoryTableRow, write_inventory_table


from core.report_contract import HEADERS, inventory_report_title


def _normalize_rows(rows: Sequence[Mapping[str, Any]]) -> list[InventoryTableRow]:
    if not isinstance(rows, list) or not rows:
        raise WDMSError("REPORT_INPUT_INVALID", "报表输入必须是非空记录列表。")

    inputs = []
    owners = set()
    for row in rows:
        if not isinstance(row, dict) or any(key not in row for key in ("owner", "whole", "loose")):
            raise WDMSError("REPORT_INPUT_INVALID", "报表记录缺少必需字段。")
        owner = row["owner"]
        if not isinstance(owner, str) or not owner.strip():
            raise WDMSError("REPORT_INPUT_INVALID", "货主必须是非空白字符串。")
        if owner in owners:
            raise WDMSError("REPORT_INPUT_INVALID", "货主名称重复。")
        owners.add(owner)

        inputs.append((owner, row["whole"], row["loose"]))

    normalized = []
    for owner, whole_value, loose_value in inputs:
        whole = parse_nonnegative_integer(whole_value, "整件数")
        loose = parse_nonnegative_quantity(loose_value, "散件数")
        normalized.append((owner, whole, loose))
    return normalized


def _check_readback(workbook: Workbook, rows: Sequence[InventoryTableRow], sheet_name: str) -> None:
    if workbook.sheetnames != [sheet_name]:
        raise WDMSError("REPORT_READBACK_MISMATCH", "报表工作表数量或名称不一致。")
    sheet = workbook[sheet_name]
    if sheet.max_row != len(rows) + 1 or sheet.max_column != 3:
        raise WDMSError("REPORT_READBACK_MISMATCH", "报表行数或列数不一致。")
    if tuple(sheet.cell(1, column).value for column in range(1, 4)) != HEADERS:
        raise WDMSError("REPORT_READBACK_MISMATCH", "报表表头不一致。")

    for index, (owner, whole, loose) in enumerate(rows, start=2):
        owner_cell = sheet.cell(index, 1)
        if owner_cell.value != owner or not isinstance(owner_cell.value, str) or owner_cell.data_type != "s":
            raise WDMSError("REPORT_READBACK_MISMATCH", "报表货主顺序或字符串类型不一致。")
        for column, expected in ((2, whole), (3, loose)):
            value = sheet.cell(index, column).value
            try:
                actual = Decimal(str(value))
                expected_decimal = Decimal(str(expected))
            except (InvalidOperation, ValueError, TypeError):
                raise WDMSError("REPORT_READBACK_MISMATCH", "报表数值回读不一致。")
            if actual != expected_decimal:
                raise WDMSError("REPORT_READBACK_MISMATCH", "报表数值回读不一致。")


def write_trial_excel(rows: list[dict], path: Union[str, Path], task_date: Optional[date] = None) -> None:
    normalized = _normalize_rows(rows)
    sheet_name = inventory_report_title(task_date if task_date is not None else date.today())
    destination = Path(path)
    temporary = destination.parent / "{}.{}.pending.xlsx".format(
        destination.stem, uuid.uuid4().hex
    )
    writer = None
    reader = None
    published = False
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        writer = Workbook()
        sheet = writer.active
        sheet.title = sheet_name
        write_inventory_table(sheet, normalized, widths=(48, 18, 22))
        writer.save(str(temporary))
        writer.close()
        writer = None

        reader = load_workbook(str(temporary), data_only=True)
        _check_readback(reader, normalized, sheet_name)
        reader.close()
        reader = None

        os.replace(str(temporary), str(destination))
        published = True
    except WDMSError:
        raise
    except Exception as error:
        raise WDMSError(
            "REPORT_PUBLISH_FAILED",
            "报表保存、打开或发布失败；原目标文件保持不变。",
        ) from error
    finally:
        pending_error = sys.exc_info()[0] is not None
        cleanup_error = None
        for workbook in (reader, writer):
            if workbook is not None:
                try:
                    workbook.close()
                except Exception as error:
                    if cleanup_error is None:
                        cleanup_error = error
        if not published and temporary.exists():
            try:
                temporary.unlink()
            except Exception as error:
                if cleanup_error is None:
                    cleanup_error = error
        if cleanup_error is not None and not pending_error:
            raise WDMSError(
                "REPORT_PUBLISH_FAILED",
                "报表临时文件清理失败。",
            ) from cleanup_error


def verify_report(report: Union[str, Path], expected_owners: Sequence[str], report_date: date) -> None:
    from openpyxl import load_workbook
    from core.numbers import parse_nonnegative_integer, parse_nonnegative_quantity

    sheet_name = inventory_report_title(report_date)
    book = None
    try:
        book = load_workbook(report, read_only=True, data_only=False)
        if book.sheetnames != [sheet_name]:
            raise ValueError("wrong sheet")
        sheet = book[sheet_name]
        if sheet.max_row != len(expected_owners) + 1 or sheet.max_column != 3:
            raise ValueError("wrong dimensions")
        rows = sheet.iter_rows()
        if tuple(cell.value for cell in next(rows)) != HEADERS:
            raise ValueError("wrong headers")
        for expected_owner, cells in zip(expected_owners, rows):
            if cells[0].value != expected_owner or cells[0].data_type != "s":
                raise ValueError("owner order differs")
            if any(cell.data_type == "f" or isinstance(cell.value, bool) for cell in cells[1:]):
                raise ValueError("unexpected formula or boolean")
            parse_nonnegative_integer(cells[1].value, "整件数")
            parse_nonnegative_quantity(cells[2].value, "散件数")
    except Exception as error:
        raise WDMSError("WORKFLOW_REPORT_UNCONFIRMED", "本轮Excel无法打开，或表头、货主顺序、行数、数值不符；未确认全流程成功。") from error
    finally:
        if book is not None:
            book.close()
