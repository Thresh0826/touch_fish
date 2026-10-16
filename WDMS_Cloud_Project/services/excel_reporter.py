"""Build, reopen, validate, and atomically publish the inventory workbook."""

from __future__ import annotations

import hashlib
import os
import uuid
from datetime import date
from pathlib import Path
from typing import Sequence, Tuple

from openpyxl import Workbook, load_workbook

from adapters.base import InventoryResult
from core.numbers import parse_nonnegative_integer, parse_nonnegative_quantity, quantity_number_format
from errors import WDMSError
from core.report_contract import HEADERS
from services.workbook_table import write_inventory_table


META_SHEET = "_WDMS_META"


class ExcelReporter:
    def publish(
        self,
        output_dir: Path,
        task_date: date,
        mode: str,
        scope: str,
        run_id: str,
        expected_count: int,
        total_owner_count: int,
        records: Sequence[InventoryResult],
    ) -> Tuple[Path, str]:
        self._validate_memory_records(records, expected_count)
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        filename = f"{task_date.year}年{task_date.month}月{task_date.day}日物流库存分析.xlsx"
        destination = output_dir / filename
        temp_path = output_dir / f".{destination.stem}.{uuid.uuid4().hex}.tmp.xlsx"
        published = False
        try:
            self._write_workbook(temp_path, task_date, mode, scope, run_id, expected_count, total_owner_count, records)
            self._validate_saved_workbook(temp_path, task_date, mode, scope, run_id, expected_count, total_owner_count, records)
            os.replace(str(temp_path), str(destination))
            published = True
            digest = sha256_file(destination)
            return destination, digest
        except WDMSError:
            raise
        except OSError as error:
            if published:
                raise WDMSError("REPORT_DIGEST_FAILED", "报表已发布，但摘要读取失败；状态不会标记完成。") from error
            raise WDMSError("REPORT_PUBLISH_FAILED", "Excel 报表保存或原子发布失败；原正式报表保持不变。") from error
        except Exception as error:
            raise WDMSError("REPORT_VALIDATION_FAILED", "Excel 报表生成或回读校验失败。") from error
        finally:
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except OSError:
                # A locked temporary file may remain; never mask the publication/validation failure.
                pass

    @staticmethod
    def _validate_memory_records(records: Sequence[InventoryResult], expected_count: int) -> None:
        if len(records) != expected_count:
            raise WDMSError("REPORT_ROW_COUNT_INVALID", "报表记录数与本次任务范围不一致。")
        names = [row.owner.name for row in records]
        if any(not name.strip() for name in names) or len(set(names)) != len(names):
            raise WDMSError("REPORT_OWNER_SET_INVALID", "报表货主名称为空或重复。")
        for row in records:
            parse_nonnegative_integer(row.whole_count, "整件数")
            parse_nonnegative_quantity(row.loose_count, "散件数")

    def _write_workbook(
        self,
        path: Path,
        task_date: date,
        mode: str,
        scope: str,
        run_id: str,
        expected_count: int,
        total_owner_count: int,
        records: Sequence[InventoryResult],
    ) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "库存分析"
        rows = [(row.owner.name, parse_nonnegative_integer(row.whole_count, "整件数"),
                 parse_nonnegative_quantity(row.loose_count, "散件数")) for row in records]
        write_inventory_table(sheet, rows, widths=(24, 14, 14), formal_style=True)

        meta = workbook.create_sheet(META_SHEET)
        meta.sheet_state = "veryHidden"
        meta.append(("run_id", run_id))
        meta.append(("date", task_date.isoformat()))
        meta.append(("mode", mode))
        meta.append(("scope", scope))
        meta.append(("expected_count", expected_count))
        meta.append(("total_owner_count", total_owner_count))
        workbook.save(path)
        workbook.close()

    def _validate_saved_workbook(
        self,
        path: Path,
        task_date: date,
        mode: str,
        scope: str,
        run_id: str,
        expected_count: int,
        total_owner_count: int,
        records: Sequence[InventoryResult],
    ) -> None:
        try:
            workbook = load_workbook(path, data_only=False, read_only=False)
        except Exception as error:
            raise WDMSError("REPORT_READBACK_FAILED", "临时报表无法重新打开。") from error
        try:
            if workbook.sheetnames != ["库存分析", META_SHEET]:
                raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表工作表结构不匹配。")
            sheet = workbook["库存分析"]
            if sheet.max_column != 3 or sheet.max_row != expected_count + 1:
                raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表行数或列数不匹配。")
            if tuple(sheet.cell(1, col).value for col in range(1, 4)) != HEADERS:
                raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表表头不匹配。")
            expected_names = [row.owner.name for row in records]
            actual_names = []
            for index, row in enumerate(records, start=2):
                name = sheet.cell(index, 1).value
                whole = sheet.cell(index, 2).value
                loose = sheet.cell(index, 3).value
                if name != row.owner.name:
                    raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表货主顺序或名称不匹配。")
                if sheet.cell(index, 1).data_type != "s":
                    raise WDMSError("REPORT_READBACK_MISMATCH", "货主名称单元格没有按文本保存。")
                if whole != row.whole_count or parse_nonnegative_quantity(loose, "散件数") != parse_nonnegative_quantity(row.loose_count, "散件数"):
                    raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表数值与内存结果不匹配。")
                if not isinstance(whole, int) or isinstance(whole, bool) or not isinstance(loose, (int, float)) or isinstance(loose, bool):
                    raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表库存字段不是有效数值单元格。")
                if sheet.cell(index, 3).number_format != quantity_number_format(row.loose_count):
                    raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表散件数的小数显示格式不匹配。")
                actual_names.append(name)
            if actual_names != expected_names or len(set(actual_names)) != expected_count:
                raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表货主集合不完整或重复。")
            if workbook[META_SHEET].sheet_state != "veryHidden":
                raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表身份元数据工作表未隐藏。")
            actual_meta = {workbook[META_SHEET].cell(row, 1).value: workbook[META_SHEET].cell(row, 2).value for row in range(1, 7)}
            expected_meta = {
                "run_id": run_id,
                "date": task_date.isoformat(),
                "mode": mode,
                "scope": scope,
                "expected_count": expected_count,
                "total_owner_count": total_owner_count,
            }
            if actual_meta != expected_meta:
                raise WDMSError("REPORT_READBACK_MISMATCH", "临时报表身份元数据与本次任务不匹配。")
        finally:
            workbook.close()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()
