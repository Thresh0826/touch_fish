"""Validate a status record and its report identity before displaying success."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

from openpyxl import load_workbook

from errors import WDMSError
from services.excel_reporter import META_SHEET, sha256_file
from services.status_manager import read_status
from services.task_lock import TaskLock


class StatusChecker:
    def __init__(self, project_root: Path, heartbeat_timeout_seconds: int = 120):
        self.project_root = Path(project_root)
        self.state_root = self.project_root / "state"
        self.lock_path = self.state_root / "task.lock"
        self.heartbeat_timeout_seconds = heartbeat_timeout_seconds

    def check(self, mode: str, today: Optional[date] = None) -> Tuple[int, Dict[str, Any]]:
        if mode not in ("mock", "real"):
            return 2, {"state": "invalid", "message": "模式参数无效。"}
        path = self.state_root / mode / "status.json"
        try:
            status = read_status(path)
        except WDMSError as error:
            return 2, {"state": "invalid", "message": error.message, "error_code": error.code}
        today = today or datetime.now().astimezone().date()
        if status is None or status.get("date") != today.isoformat():
            return 2, {"state": "missing", "message": "未检测到今日有效任务结果。"}
        if status.get("mode") != mode:
            return 2, {"state": "invalid", "message": "状态文件模式与查询模式不一致。"}

        state = status.get("status")
        if state == "running":
            lock_held = TaskLock.is_locked(self.lock_path)
            heartbeat = self._parse_time(status.get("heartbeat_time"))
            if not lock_held:
                return 2, {"state": "interrupted", "message": "任务异常中断，状态不可确认。"}
            if heartbeat is None or (datetime.now().astimezone() - heartbeat).total_seconds() > self.heartbeat_timeout_seconds:
                return 2, {"state": "uncertain", "message": "任务心跳已过期，状态不可确认。"}
            return 1, {"state": "running", "message": "今日库存采集任务仍在执行。"}
        if state == "failed":
            reason = status.get("error") or "任务未能完成。"
            return 2, {
                "state": "failed",
                "error_code": status.get("error_code"),
                "message": f"今日库存采集失败：{reason}",
            }
        if state != "completed_pending_review":
            return 2, {"state": "invalid", "message": "状态值无效，不能认定任务完成。"}

        mismatch = self._validate_report(status, mode)
        if mismatch:
            return 2, {"state": "invalid", "message": mismatch}
        if mode == "mock":
            if status.get("scope") != "mock":
                return 2, {"state": "invalid", "message": "模拟状态范围不匹配。"}
            if status.get("expected_count") != 24 or status.get("completed_count") != 24 or status.get("total_owner_count") != 24:
                return 2, {"state": "invalid", "message": "模拟货主数量校验失败。"}
            return 0, {"state": "mock_complete", "message": "模拟报表已生成，结果仅供测试。"}
        if status.get("scope") == "trial":
            return 1, {"state": "trial_pending_review", "message": "Real 试验结果已生成，等待现场核验。"}
        if status.get("scope") != "full":
            return 2, {"state": "invalid", "message": "Real 状态范围不是 24 个货主全量任务。"}
        if status.get("expected_count") != 24 or status.get("completed_count") != 24 or status.get("total_owner_count") != 24:
            return 2, {"state": "invalid", "message": "Real 全量货主数量校验失败。"}
        return 0, {"state": "completed_pending_review", "message": "今日库存分析报表已经生成，等待人工核验。"}

    def _validate_report(self, status: Dict[str, Any], mode: str) -> Optional[str]:
        output_path = status.get("output_file")
        if not output_path:
            return "状态未关联本次报表，不能认定任务完成。"
        path = Path(output_path)
        if not path.is_file():
            return "本次报表文件缺失，不能认定任务完成。"
        try:
            if sha256_file(path) != status.get("sha256"):
                return "报表摘要与状态记录不一致，不能认定任务完成。"
            workbook = load_workbook(path, read_only=True, data_only=False)
        except Exception:
            return "报表无法读取或摘要校验失败。"
        try:
            if META_SHEET not in workbook.sheetnames:
                return "报表缺少本次任务身份信息。"
            sheet = workbook[META_SHEET]
            metadata = {sheet.cell(row, 1).value: sheet.cell(row, 2).value for row in range(1, 7)}
            expected = {
                "run_id": status.get("run_id"),
                "date": status.get("date"),
                "mode": mode,
                "scope": status.get("scope"),
                "expected_count": status.get("expected_count"),
                "total_owner_count": status.get("total_owner_count"),
            }
            if metadata != expected:
                return "报表身份信息与状态记录不一致。"
        finally:
            workbook.close()
        if status.get("completed_count") != status.get("expected_count"):
            return "已完成货主数与预期数量不一致。"
        return None

    @staticmethod
    def _parse_time(value: Any) -> Optional[datetime]:
        if not isinstance(value, str):
            return None
        try:
            parsed = datetime.fromisoformat(value)
            return parsed if parsed.tzinfo else None
        except ValueError:
            return None
