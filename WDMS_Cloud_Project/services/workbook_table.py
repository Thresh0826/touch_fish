"""Shared typed inventory table; report identities stay with each publisher."""
from __future__ import annotations

from decimal import Decimal
from typing import Sequence, Tuple

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

from core.numbers import quantity_number_format
from core.report_contract import HEADERS

InventoryTableRow = Tuple[str, int, Decimal]


def write_inventory_table(sheet: Worksheet, rows: Sequence[InventoryTableRow], *,
                          widths: Tuple[int, int, int], formal_style: bool = False) -> None:
    """Preserve decimal cells and formula-like owner names as literal text."""
    sheet.append(HEADERS)
    for owner, whole, loose in rows:
        sheet.append((None, whole, loose))
        owner_cell = sheet.cell(sheet.max_row, 1)
        owner_cell.value = owner
        owner_cell.data_type = 's'
        sheet.cell(sheet.max_row, 2).number_format = '#,##0'
        sheet.cell(sheet.max_row, 3).number_format = quantity_number_format(loose)
    for column, width in zip(('A', 'B', 'C'), widths):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    if formal_style:
        header_font = Font(name='Arial', bold=True, color='000000')
        body_font = Font(name='Arial', color='000000')
        for cell in sheet[1]:
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(bottom=Side(style='thin', color='000000'))
        for cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=3):
            for cell in cells:
                cell.font = body_font
            cells[0].alignment = Alignment(horizontal='left')
            cells[1].alignment = Alignment(horizontal='right')
            cells[2].alignment = Alignment(horizontal='right')
        sheet.auto_filter.ref = f'A1:C{sheet.max_row}'
        sheet.sheet_view.showGridLines = False
