"""Read-only XLSX input and immutable per-row queries; no live spreadsheet control."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
import hashlib
from io import BytesIO
import math
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

from moyu_tools.model import fail

MAX_TASKS = 10000
MAX_SCAN_ROWS = 100000
MAX_COLUMNS = 256
IDENTIFIER = re.compile(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}")


def identifier(value, number_format="General"):
    if isinstance(value, str):
        result = value.strip()
    elif type(value) in (int, float):
        if not math.isfinite(value) or value < 0 or value != int(value) or value >= 10**15:
            raise ValueError("单据编号必须为完整文本或不超过15位的非负整数")
        result = str(int(value))
        if re.fullmatch(r"0{1,80}", number_format):
            result = result.zfill(len(number_format))
        elif number_format not in ("General", "@"):
            raise ValueError("单据编号的数字格式不明确，请在源表中改为完整文本")
    else:
        raise ValueError("单据编号不能为空、公式或布尔值")
    if not IDENTIFIER.fullmatch(result):
        raise ValueError("单据编号须为1至80位的字母、数字、下划线或短横线")
    return result


def calendar_date(value):
    if type(value) is datetime:
        return value.date()
    if type(value) is date:
        return value
    if isinstance(value, str):
        match = re.fullmatch(r"([0-9]{4})[-/.年]([0-9]{1,2})[-/.月]([0-9]{1,2})日?", value.strip())
        if match:
            return date(*(int(part) for part in match.groups()))
    raise ValueError("下单日期须为Excel日期或年月日文本，例如2023-12-18")


@dataclass(frozen=True)
class SheetRow:
    source_row: int
    day: date
    contract: str

    def __post_init__(self):
        if (type(self.source_row) is not int or not 1 <= self.source_row <= MAX_SCAN_ROWS
                or type(self.day) is not date or not isinstance(self.contract, str)
                or not IDENTIFIER.fullmatch(self.contract)):
            fail("MOYU_SHEET_INVALID", "表格任务行无效，未开始查询。")

    @property
    def start(self):
        return self.day

    @property
    def end(self):
        return self.day

    @property
    def key(self):
        return f"row_{self.source_row:06d}_{self.day.isoformat()}"

    def document(self):
        return {"source_row": self.source_row, "date": self.day.isoformat(), "contract": self.contract}


@dataclass(frozen=True)
class SheetPlan:
    source_name: str
    source_sha256: str
    sheet: str
    header_row: int
    rows: tuple

    def __post_init__(self):
        if (not isinstance(self.source_name, str) or not self.source_name
                or len(self.source_name) > 255 or not isinstance(self.source_sha256, str)
                or not re.fullmatch(r"[a-f0-9]{64}", self.source_sha256)
                or not isinstance(self.sheet, str) or not self.sheet or len(self.sheet) > 31
                or type(self.header_row) is not int or not 1 <= self.header_row <= 30
                or type(self.rows) is not tuple or not 1 <= len(self.rows) <= MAX_TASKS
                or any(not isinstance(row, SheetRow) for row in self.rows)
                or any(row.source_row <= self.header_row for row in self.rows)
                or [row.source_row for row in self.rows] != sorted({row.source_row for row in self.rows})):
            fail("MOYU_SHEET_INVALID", "表格任务计划无效，请重新导入；旧进度保留。")

    def segments(self):
        return iter(self.rows)

    def document(self):
        return {"kind": "sheet_rows", "source_name": self.source_name, "source_sha256": self.source_sha256,
                "sheet": self.sheet, "header_row": self.header_row, "rows": [row.document() for row in self.rows]}

    @classmethod
    def parse(cls, data):
        try:
            return cls(data["source_name"], data["source_sha256"], data["sheet"], data["header_row"],
                       tuple(SheetRow(row["source_row"], date.fromisoformat(row["date"]), row["contract"])
                             for row in data["rows"]))
        except (KeyError, TypeError, ValueError):
            fail("MOYU_SHEET_INVALID", "保存的表格计划不完整，请保留旧任务并重新导入。")


def source_bytes(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        fail("MOYU_XLSX_INVALID", "请选择不超过32MB的.xlsx文件，原表不会被修改。")
    with path.open("rb") as source:
        data = source.read(32*1024*1024+1)
    if len(data) > 32*1024*1024:
        fail("MOYU_XLSX_INVALID", "请选择不超过32MB的.xlsx文件，原表不会被修改。")
    return data


def workbook(path):
    """Open bounded local OOXML, retaining formulas so they cannot become identifiers."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        fail("MOYU_XLSX_DEPENDENCY", "当前Python缺少openpyxl，请选择原WDMS已配置的Python环境；工具不会联网安装。")
    data = source_bytes(path)
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 128*1024*1024:
                fail("MOYU_XLSX_INVALID", "表格解压后过大，请另存本次需要的工作表。")
        return (load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False),
                hashlib.sha256(data).hexdigest())
    except (OSError, ValueError, BadZipFile, KeyError):
        fail("MOYU_XLSX_INVALID", "无法读取.xlsx，请在WPS或Excel中另存为普通.xlsx后重试。")


def sheet_names(path):
    book, _ = workbook(path)
    try:
        return [sheet.title for sheet in book.worksheets if sheet.sheet_state == "visible"]
    finally:
        book.close()


def import_sheet(path, sheet):
    path = Path(path)
    book, before = workbook(path)
    try:
        if sheet not in book.sheetnames or book[sheet].sheet_state != "visible":
            fail("MOYU_SHEET_INVALID", "请选择一个可见工作表。")
        worksheet = book[sheet]
        # Do not trust a producer's used-range metadata; blank formatting far below
        # the input does not create jobs, and underestimated dimensions cannot drop rows.
        worksheet.reset_dimensions()
        headers, jobs = None, []
        for number, cells in enumerate(worksheet.iter_rows(max_col=MAX_COLUMNS+1), 1):
            if number > MAX_SCAN_ROWS:
                fail("MOYU_SHEET_INVALID", "工作表超过10万行，请另存本次需要处理的行。")
            if cells[MAX_COLUMNS].value is not None:
                fail("MOYU_SHEET_INVALID", "工作表超过256列，请另存本次需要的列。")
            labels = [str(cell.value).strip() if cell.value is not None else "" for cell in cells]
            if number <= 30 and "下单日期" in labels and "单据编号" in labels:
                if headers is not None or labels.count("下单日期") != 1 or labels.count("单据编号") != 1:
                    fail("MOYU_SHEET_HEADERS", "存在重复的表头，请保留一组“下单日期”和“单据编号”。")
                headers = (number, labels.index("下单日期"), labels.index("单据编号"))
                continue
            if headers is None:
                if number >= 30:
                    fail("MOYU_SHEET_HEADERS", "前30行没有找到“下单日期”和“单据编号”表头。")
                continue
            _, date_column, contract_column = headers
            date_cell, contract_cell = cells[date_column], cells[contract_column]
            present = lambda cell: cell.value is not None and (not isinstance(cell.value, str) or bool(cell.value.strip()))
            if not present(date_cell) and not present(contract_cell):
                continue
            if date_cell.data_type in ("f", "e") or contract_cell.data_type in ("f", "e"):
                fail("MOYU_SHEET_ROW_INVALID", f"工作表第{number}行含公式或错误值，请把日期和单据编号保存为实际值。")
            try:
                jobs.append(SheetRow(number, calendar_date(date_cell.value),
                                     identifier(contract_cell.value, contract_cell.number_format)))
            except ValueError as error:
                fail("MOYU_SHEET_ROW_INVALID", f"工作表第{number}行：{error}；整表未开始运行。")
            if len(jobs) > MAX_TASKS:
                fail("MOYU_SHEET_INVALID", "本次最多导入1万条记录，请拆成多个任务。")
        if headers is None:
            fail("MOYU_SHEET_HEADERS", "没有找到“下单日期”和“单据编号”表头。")
        if not jobs:
            fail("MOYU_SHEET_INVALID", "所选工作表没有可处理的日期和单据编号。")
    finally:
        book.close()
    after = hashlib.sha256(source_bytes(path)).hexdigest()
    if before != after:
        fail("MOYU_SHEET_CHANGED", "导入期间表格发生变化，请保存表格后重新导入。")
    return SheetPlan(path.name, after, sheet, headers[0], tuple(jobs))
