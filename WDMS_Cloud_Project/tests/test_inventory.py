from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import time
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from adapters.base import InventoryResult, Owner
from adapters.mock_wdms import MOCK_RECORDS, MockWDMSAdapter
from adapters.supervisor import JsonLineWorker
from adapters.worker import WindowsWorker
from config import validate_common_config
from core.inventory_service import InventoryService
from core.numbers import parse_nonnegative_integer
from core.task_runner import TaskRunner
from errors import WDMSError
from services.excel_reporter import ExcelReporter
from services.status_checker import StatusChecker
from services.task_lock import TaskLock


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORK_ROOT = Path(tempfile.gettempdir())


def _temp_project(config_changes=None):
    temp = tempfile.TemporaryDirectory(dir=WORK_ROOT)
    root = Path(temp.name)
    with (PROJECT_ROOT / "config.sample.json").open("r", encoding="utf-8") as source:
        config = json.load(source)
    if config_changes:
        config.update(config_changes)
    config_path = root / "config.json"
    config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    return temp, root, config_path


class TestInventoryFlow(unittest.TestCase):
    def test_mock_full_workbook_and_status(self):
        temp, root, config_path = _temp_project()
        self.addCleanup(temp.cleanup)
        output, status_path = TaskRunner(root, config_path).run("mock")
        workbook = load_workbook(output, data_only=False)
        try:
            sheet = workbook["库存分析"]
            self.assertEqual(workbook.sheetnames, ["库存分析", "_WDMS_META"])
            self.assertEqual((sheet.max_row, sheet.max_column), (25, 3))
            self.assertEqual([sheet.cell(1, col).value for col in range(1, 4)], ["货主", "整件数", "散件数"])
            self.assertEqual([sheet.cell(2, col).value for col in range(1, 4)], ["模拟货主01", 0, 0])
            self.assertEqual([sheet.cell(3, col).value for col in (2, 3)], [sheet.cell(4, col).value for col in (2, 3)])
            self.assertEqual(sheet.cell(25, 1).value, "模拟货主24")
            self.assertEqual(sheet.cell(2, 1).data_type, "s")
            self.assertEqual(workbook["_WDMS_META"].sheet_state, "veryHidden")
        finally:
            workbook.close()
        status = json.loads(status_path.read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "completed_pending_review")
        self.assertEqual((status["completed_count"], status["expected_count"], status["total_owner_count"]), (24, 24, 24))
        self.assertNotIn("inventory", json.dumps(status).lower())
        code, result = StatusChecker(root).check("mock")
        self.assertEqual(code, 0)
        self.assertEqual(result["state"], "mock_complete")

    def test_stale_or_wrong_result_fails_without_report(self):
        for fault in (
            {"stale_result": ["模拟货主02"]},
            {"wrong_owner": ["模拟货主02"]},
            {"false_complete": ["模拟货主02"]},
            {"mutate_during_read": ["模拟货主02"]},
            {"timeout": ["模拟货主02"]},
        ):
            temp, root, config_path = _temp_project({"mock_faults": fault})
            try:
                with self.assertRaises(WDMSError):
                    TaskRunner(root, config_path).run("mock")
                status = json.loads((root / "state/mock/status.json").read_text(encoding="utf-8"))
                self.assertEqual(status["status"], "failed")
                self.assertIsNone(status["output_file"])
                self.assertFalse(list((root / "output/mock").glob("*.xlsx")))
            finally:
                temp.cleanup()

    def test_incomplete_roster_and_read_error_fail(self):
        expected = [Owner(name, name) for name, _, _ in MOCK_RECORDS]
        for fault in ({"duplicate_owner_list": True}, {"missing_owner_list": True}):
            adapter = MockWDMSAdapter(fault)
            adapter.connect()
            with self.assertRaises(WDMSError):
                InventoryService(adapter).collect(expected)
            adapter.close()
        adapter = MockWDMSAdapter({"read_fail": ["模拟货主02"]})
        with self.assertRaises(WDMSError):
            InventoryService(adapter).collect(expected)
        adapter.close()

    def test_count_parser_preserves_zero_and_rejects_bad_values(self):
        self.assertEqual(parse_nonnegative_integer("0", "整件数"), 0)
        self.assertEqual(parse_nonnegative_integer("1,234", "整件数"), 1234)
        for bad in ("", "  ", "-1", "1.2", "3,12", "12 件", 1.0, True):
            with self.subTest(bad=bad):
                with self.assertRaises(WDMSError):
                    parse_nonnegative_integer(bad, "整件数")
        with self.assertRaises(WDMSError) as context:
            parse_nonnegative_integer("1000000000000000", "整件数")
        self.assertEqual(context.exception.code, "COUNT_PRECISION_RISK")

    def test_old_date_is_not_valid_today(self):
        temp, root, config_path = _temp_project()
        self.addCleanup(temp.cleanup)
        TaskRunner(root, config_path).run("mock")
        code, result = StatusChecker(root).check("mock", today=date.today() + timedelta(days=1))
        self.assertEqual(code, 2)
        self.assertEqual(result["state"], "missing")

    def test_report_save_failure_preserves_existing_file(self):
        temp = tempfile.TemporaryDirectory(dir=WORK_ROOT)
        self.addCleanup(temp.cleanup)
        output_dir = Path(temp.name)
        destination = output_dir / f"{date.today().year}年{date.today().month}月{date.today().day}日物流库存分析.xlsx"
        destination.write_bytes(b"previous valid report")
        adapter = MockWDMSAdapter()
        adapter.connect()
        owners = list(adapter.list_owners())
        records = InventoryService(adapter).collect(owners)
        adapter.close()
        with mock.patch.object(ExcelReporter, "_write_workbook", side_effect=OSError("simulated disk failure")):
            with self.assertRaises(WDMSError) as context:
                ExcelReporter().publish(output_dir, date.today(), "mock", "mock", "run-test", 24, 24, records)
        self.assertEqual(context.exception.code, "REPORT_PUBLISH_FAILED")
        self.assertEqual(destination.read_bytes(), b"previous valid report")

    def test_digest_failure_reports_that_new_report_was_already_published(self):
        temp = tempfile.TemporaryDirectory(dir=WORK_ROOT)
        self.addCleanup(temp.cleanup)
        output_dir = Path(temp.name)
        destination = output_dir / f"{date.today().year}年{date.today().month}月{date.today().day}日物流库存分析.xlsx"
        destination.write_bytes(b"previous report")
        adapter = MockWDMSAdapter()
        adapter.connect()
        owners = list(adapter.list_owners())
        records = InventoryService(adapter).collect(owners)
        adapter.close()
        with mock.patch("services.excel_reporter.sha256_file", side_effect=OSError("simulated read failure")):
            with self.assertRaises(WDMSError) as context:
                ExcelReporter().publish(output_dir, date.today(), "mock", "mock", "run-digest", 24, 24, records)
        self.assertEqual(context.exception.code, "REPORT_DIGEST_FAILED")
        workbook = load_workbook(destination, read_only=True)
        try:
            self.assertEqual(workbook["库存分析"].max_row, 25)
        finally:
            workbook.close()

    def test_formula_like_owner_name_is_written_as_literal_text(self):
        temp = tempfile.TemporaryDirectory(dir=WORK_ROOT)
        self.addCleanup(temp.cleanup)
        record = InventoryResult("req", Owner("test", "=1+1"), 0, 8, {"test": True})
        output, _ = ExcelReporter().publish(Path(temp.name), date.today(), "mock", "mock", "run-text", 1, 24, [record])
        workbook = load_workbook(output, data_only=False)
        try:
            cell = workbook["库存分析"]["A2"]
            self.assertEqual(cell.value, "=1+1")
            self.assertEqual(cell.data_type, "s")
        finally:
            workbook.close()

    def test_non_finite_timeout_configuration_is_rejected(self):
        for value in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(value=value):
                with self.assertRaises(WDMSError):
                    validate_common_config({"timing": {"query_timeout_seconds": value}})

    def test_keyboard_interrupt_commits_failed_status_and_releases_lock(self):
        temp, root, config_path = _temp_project()
        self.addCleanup(temp.cleanup)
        with mock.patch("core.task_runner.InventoryService.collect", side_effect=KeyboardInterrupt()):
            with self.assertRaises(WDMSError) as context:
                TaskRunner(root, config_path).run("mock")
        self.assertEqual(context.exception.code, "TASK_INTERRUPTED")
        status = json.loads((root / "state/mock/status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "failed")
        self.assertIsNone(status["output_file"])
        self.assertFalse(TaskLock.is_locked(root / "state/task.lock"))

    def test_cross_process_lock_is_shared_independent_of_output_mode(self):
        temp = tempfile.TemporaryDirectory(dir=WORK_ROOT)
        self.addCleanup(temp.cleanup)
        lock_path = Path(temp.name) / "state/task.lock"
        lock = TaskLock(lock_path)
        lock.acquire()
        try:
            script = (
                "import sys\nfrom pathlib import Path\n"
                "from services.task_lock import TaskLock\nfrom errors import TaskAlreadyRunning\n"
                "lock = TaskLock(Path(sys.argv[1]))\n"
                "try:\n    lock.acquire()\n"
                "except TaskAlreadyRunning:\n    print('busy')\n"
                "else:\n    lock.release()\n    print('acquired')\n"
            )
            result = subprocess.run([sys.executable, "-c", script, str(lock_path)], cwd=PROJECT_ROOT, capture_output=True, text=True, timeout=5)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(result.stdout.strip(), "busy")
        finally:
            lock.release()
        result = subprocess.run([sys.executable, "-c", script, str(lock_path)], cwd=PROJECT_ROOT, capture_output=True, text=True, timeout=5)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), "acquired")

    def test_supervisor_kills_blocked_worker_with_hard_timeout(self):
        temp = tempfile.TemporaryDirectory(dir=WORK_ROOT)
        self.addCleanup(temp.cleanup)
        script = Path(temp.name) / "blocking_worker.py"
        script.write_text("import sys, time\nsys.stdin.readline()\ntime.sleep(20)\n", encoding="utf-8")
        worker = JsonLineWorker([sys.executable, str(script)], cwd=temp.name)
        start = time.monotonic()
        with self.assertRaises(WDMSError) as context:
            worker.request("block", {}, 0.25)
        elapsed = time.monotonic() - start
        self.assertEqual(context.exception.code, "WORKER_HARD_TIMEOUT")
        self.assertLess(elapsed, 2.0)
        self.assertIsNotNone(worker.process.poll())

    @unittest.skipIf(sys.platform == "win32", "macOS-only boundary check")
    def test_real_mode_fails_before_loading_windows_automation(self):
        temp, root, config_path = _temp_project()
        self.addCleanup(temp.cleanup)
        with self.assertRaises(WDMSError) as context:
            TaskRunner(root, config_path).run("real", limit=1)
        self.assertEqual(context.exception.code, "REAL_MODE_WINDOWS_ONLY")
        self.assertNotIn("pywinauto", sys.modules)
        status = json.loads((root / "state/real/status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "failed")


class FakeCombo:
    def __init__(self, value="", refuse=False):
        self.value = value
        self.refuse = refuse

    def select(self, value):
        if not self.refuse:
            self.value = value

    def selected_text(self):
        return self.value


class FakeControl:
    def __init__(self, value=""):
        self.value = value
        self.calls = 0
        self.inputs = []

    def invoke(self):
        self.calls += 1

    def set_edit_text(self, value):
        self.inputs.append(value)

    def window_text(self):
        if isinstance(self.value, list):
            index = min(self.calls, len(self.value) - 1)
            self.calls += 1
            return self.value[index]
        return self.value


class TestRealWorkerLogicWithFakeControls(unittest.TestCase):
    def _make_worker(self):
        evidence = {
            "strategy": "version",
            "confirmed": True,
            "validation_record": "TEST-ONLY",
            "version_control_id": "version",
            "version_text_source": "window_text",
            "complete_control_id": "complete",
            "complete_text_source": "window_text",
            "complete_value": "完成",
            "result_owner_control_id": "result_owner",
            "result_owner_text_source": "window_text",
        }
        config = {"timing": {"query_timeout_seconds": 2}, "real": {"main_window": {}, "login": {}, "navigation": {"steps": []}, "controls": {
            "page_container_id": "page", "owner_combo_id": "cmbLjq", "search_button_id": "btnSearch",
            "whole_count_id": "lblZjs", "loose_count_id": "lblSjs",
        }, "query_evidence": evidence}}
        worker = WindowsWorker(config)
        worker._require_inventory_controls = lambda: None
        worker._preflight = lambda **kwargs: None
        worker.controls = {
            "cmbLjq": FakeCombo("貨主甲"),
            "btnSearch": FakeControl(),
            "version": FakeControl(["7", "8"]),
            "complete": FakeControl("完成"),
            "result_owner": FakeControl("貨主甲"),
            "lblZjs": FakeControl("0"),
            "lblSjs": FakeControl("8"),
        }
        return worker

    def test_exact_owner_select_invoke_and_version_bound_read(self):
        worker = self._make_worker()
        owner = Owner("甲", "貨主甲")
        worker.select_owner({"owner_id": owner.owner_id, "owner_name": owner.name})
        result = worker.start_query({"owner_id": owner.owner_id, "owner_name": owner.name})
        self.assertEqual(worker.controls["btnSearch"].calls, 1)
        actual = worker.read_result({"request_id": result["request_id"], "owner_id": owner.owner_id, "owner_name": owner.name, "timeout_seconds": 2})
        self.assertEqual((actual["whole_count"], actual["loose_count"]), ("0", "8"))
        self.assertTrue(actual["evidence"]["owner_match"])

    def test_selection_failure_does_not_trigger_query(self):
        worker = self._make_worker()
        worker.controls["cmbLjq"] = FakeCombo("其他货主", refuse=True)
        with self.assertRaises(WDMSError) as context:
            worker.select_owner({"owner_id": "甲", "owner_name": "貨主甲"})
        self.assertEqual(context.exception.code, "OWNER_SELECTION_FAILED")

    def test_navigation_invokes_only_configured_steps_and_checks_markers(self):
        worker = self._make_worker()
        worker._require_connected = lambda: None
        steps = []
        controls = {}
        for index in range(1, 4):
            step_id = f"menu{index}"
            marker_id = f"page{index}"
            controls[step_id] = FakeControl()
            controls[marker_id] = FakeControl(f"目标页{index}")
            steps.append({"label": f"步骤{index}", "automation_id": step_id, "page_marker_id": marker_id, "page_marker_value": f"目标页{index}", "action": "invoke", "timeout_seconds": 1})
        worker.real["navigation"] = {"steps": steps}
        worker.main = object()
        worker._find_unique = lambda root, automation_id, required=True: controls.get(automation_id)
        worker._cache_inventory_controls = lambda: None
        self.assertEqual(worker.navigate({}), {"navigated": True})
        self.assertEqual([controls[f"menu{i}"].calls for i in range(1, 4)], [1, 1, 1])

    def test_already_logged_in_branch_requires_exact_session_marker(self):
        worker = WindowsWorker({"real": {"main_window": {}, "login": {
            "window_title_regex": "登录", "session_marker_id": "session", "session_marker_value": "已登录",
        }, "controls": {}, "query_evidence": {}}})
        worker.main = object()
        worker.pid = 123
        worker._require_connected = lambda: None
        worker._preflight = lambda **kwargs: None
        worker._login_windows = lambda: []
        marker = FakeControl("已登录")
        worker._find_unique = lambda root, automation_id, required=True: marker
        self.assertEqual(worker.login({}), {"logged_in": True, "branch": "existing_session"})

    def test_login_credentials_branch_uses_configured_controls_and_markers(self):
        login_cfg = {
            "username": "test-user", "credential_target": "test-target", "window_title_regex": "登录",
            "session_marker_id": "session", "session_marker_value": "已登录",
            "success_marker_id": "success", "success_marker_value": "已登录",
            "failure_marker_id": "failure", "failure_marker_value": "登录失败",
            "username_control_id": "user", "password_control_id": "password", "submit_button_id": "submit",
        }
        worker = WindowsWorker({"real": {"main_window": {}, "login": login_cfg, "navigation": {"steps": []}, "controls": {}, "query_evidence": {}}})
        worker.main = object()
        worker.pid = 123
        worker._require_connected = lambda: None
        worker._preflight = lambda **kwargs: None
        worker._login_windows = lambda: [42]
        worker._uia_window = lambda hwnd: object()
        user = FakeControl()
        password = FakeControl()
        button = FakeControl()
        failure = FakeControl("")
        success = FakeControl("已登录")
        controls = {"user": user, "password": password, "submit": button, "failure": failure}
        worker._find_controls = lambda root, ids: controls
        worker._read_credential = lambda: "test-only-secret"
        worker._find_unique = lambda root, automation_id, required=True: success
        self.assertEqual(worker.login({}), {"logged_in": True, "branch": "credentials"})
        self.assertEqual(user.inputs, ["test-user"])
        self.assertEqual(password.inputs, ["test-only-secret"])
        self.assertEqual(button.calls, 1)


if __name__ == "__main__":
    unittest.main()
