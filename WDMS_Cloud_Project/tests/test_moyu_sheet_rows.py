"""Synthetic XLSX and per-document loops; no real workbook or Windows access."""
from dataclasses import asdict
from datetime import date, datetime
import hashlib
from pathlib import Path
import queue
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from core.step_executor import StepContext
from errors import WDMSError
from moyu_tools.engine import Engine
from moyu_tools.bridge import Bridge
from moyu_tools.model import Plan, Policy
from moyu_tools.native import LocalQuery
from moyu_tools.runtime import make_request, verify_child, worker
from moyu_tools.settings import inventory_lock_path, load_json, state_root, task_lock_path, validate_profile
from moyu_tools.sheet import SheetPlan, SheetRow, calendar_date, identifier, import_sheet, sheet_names
from moyu_tools.store import Store
from services.status_manager import atomic_json_write
from services.task_lock import TaskLock
from tests.moyu_fixtures import Clock, Remote, environment, profile, sheet_profile, zip_bytes


def row_plan():
    return SheetPlan("synthetic.xlsx", "a"*64, "任务", 1, (
        SheetRow(2, date(2024, 3, 2), "000041"),
        SheetRow(3, date(2024, 2, 29), "000042"),
        SheetRow(5, date(2024, 3, 2), "000043")))


class Page:
    """Native-read/write protocol only. All cursor/OS behaviour remains unverified."""
    def __init__(self, clock):
        self.clock, self.bridge, self.api = clock, self, Mock()
        self.log = None
        self.profile = sheet_profile()
        self.values = {"start": "2024-01-01", "end": "2024-01-01", "contract": "old"}
        self.results = {"000041": "001111", "000042": None, "000043": "002222"}
        self.started, self.clicks, self.actions = None, [], []
        self.label, self.fixed_hash = "合同编号", "b"*64
        self.changed_readback = False
        self.api.click_query_input.side_effect = self.click

    def control(self, target, root, context):
        context.check()

    def guard(self, root, context):
        context.check()

    def click(self, root, hwnd, log):
        self.api.query_input_guard()
        self.clicks.append(dict(self.values))
        self.started = self.clock()

    def signal(self, segment, context):
        context.check()
        return {"busy": self.started is not None and self.clock()-self.started < 1, "responsive": True}

    def call(self, action, payload, context):
        context.check()
        self.actions.append((action, dict(payload)))
        if action == "conditions":
            return {"hash": self.fixed_hash}
        if action in ("date", "edit"):
            key = ("start" if payload["hwnd"] == self.profile["wdms"]["start"]["hwnd"] else "end") if action == "date" else "contract"
            if "value" in payload:
                self.values[key] = payload["value"]
            actual = self.values[key]
            if action == "edit" and self.changed_readback and "value" not in payload:
                actual = "different"
            return {"value": actual}
        if action == "read":
            key = next(key for key in ("query", "contract_label", "count_label", "count", "header", "cell")
                       if self.profile["wdms"][key]["point"] == payload["point"])
            number = self.results.get(self.values["contract"])
            if key == "header":
                return {"name": "送货单号", "role": 25, "rect": [600, 330, 800, 350]}
            if key == "cell":
                return {"value": number, "role": 29, "rect": [600, 350, 800, 375]}
            return {"name": {"query": "查询(F)", "contract_label": self.label,
                             "count_label": "总行数", "count": "1" if number else "0"}[key]}
        raise AssertionError(action)

    def query(self):
        query = LocalQuery(self, self.profile)
        query.attach = lambda context: context.check()
        query.signal = self.signal
        return query


class NamedRemote(Remote):
    def search(self, number, context):
        self.name, self.data = "actual_"+number+".zip", zip_bytes(number.encode("ascii"))
        return super().search(number, context)


class MoyuSheetImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name)/"synthetic.xlsx"

    def book(self, rows, headers=("货位类型", "下单日期", "单据编号", "仓库"), *, title_rows=0, epoch=None):
        book = Workbook()
        sheet = book.active
        sheet.title = "任务"
        if epoch:
            book.epoch = epoch
        for _ in range(title_rows):
            sheet.append(["合成测试表，不含真实业务数据"])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        book.save(self.path)
        book.close()
        return book

    def test_original_order_leap_day_leading_zeros_and_same_day_documents(self):
        self.book([("合成", datetime(2024, 3, 2), "000041", "测试仓"),
                   ("合成", "2024/2/29", 42, "测试仓"),
                   (None, None, None, None),
                   ("合成", "2024-03-02", "000043", "测试仓")])
        before = self.path.read_bytes()
        plan = import_sheet(self.path, "任务")
        self.assertEqual([row.source_row for row in plan.rows], [2, 3, 5])
        self.assertEqual([row.contract for row in plan.rows], ["000041", "42", "000043"])
        self.assertEqual([row.start for row in plan.rows], [date(2024, 3, 2), date(2024, 2, 29), date(2024, 3, 2)])
        self.assertTrue(all(row.start == row.end for row in plan.rows))
        self.assertEqual(plan.source_sha256, hashlib.sha256(before).hexdigest())
        self.assertEqual(self.path.read_bytes(), before)

    def test_formatted_numeric_identifier_retains_displayed_zero_padding(self):
        book = self.book([(None, "2024-02-29", 41, None)])
        book["任务"]["C2"].number_format = "000000"
        book.save(self.path)
        book.close()
        self.assertEqual(import_sheet(self.path, "任务").rows[0].contract, "000041")

    def test_sheet_choice_and_header_position_are_explicit(self):
        book = self.book([(None, "2024.2.29", "000041", None)], title_rows=2)
        book.create_sheet("隐藏").sheet_state = "hidden"
        other = book.create_sheet("另一个表")
        other.append(["下单日期", "单据编号"])
        other.append(["2025-01-01", "000099"])
        book.save(self.path)
        book.close()
        self.assertEqual(sheet_names(self.path), ["任务", "另一个表"])
        self.assertEqual(import_sheet(self.path, "任务").header_row, 3)
        self.assertEqual(import_sheet(self.path, "另一个表").rows[0].contract, "000099")
        with self.assertRaises(WDMSError):
            import_sheet(self.path, "隐藏")

    def test_1904_epoch_dates_are_read_as_their_actual_calendar_day(self):
        self.book([(None, datetime(2024, 2, 29), "000041", None)], epoch=CALENDAR_MAC_1904)
        self.assertEqual(import_sheet(self.path, "任务").rows[0].day, date(2024, 2, 29))

    def test_hidden_and_filtered_rows_are_included_without_merging_duplicates(self):
        book = self.book([(None, "2024-02-29", "000041", None)]*2)
        book["任务"].row_dimensions[3].hidden = True
        book["任务"].auto_filter.ref = "A1:D3"
        book.save(self.path)
        book.close()
        plan = import_sheet(self.path, "任务")
        self.assertEqual(len(plan.rows), 2)
        self.assertNotEqual(plan.rows[0].key, plan.rows[1].key)

    def test_formula_error_partial_and_invalid_date_abort_the_entire_import(self):
        for day, number in (("2024-02-30", "000041"), (None, "000041"), ("2024-02-29", None),
                            ("=TODAY()", "000041"), ("2024-02-29", '=TEXT(41,"000000")'),
                            ("2024-02-29", "#N/A"), (45000, "000041")):
            with self.subTest(day=day, number=number):
                self.book([(None, "2024-02-29", "000001", None), (None, day, number, None)])
                with self.assertRaises(WDMSError) as caught:
                    import_sheet(self.path, "任务")
                self.assertEqual(caught.exception.code, "MOYU_SHEET_ROW_INVALID")
                self.assertIn("3", caught.exception.message)

    def test_missing_duplicate_and_late_headers_are_not_guessed(self):
        for headers, offset in ((("日期", "单据编号"), 0), (("下单日期", "单据编号", "单据编号"), 0),
                                (("下单日期", "单据编号"), 30)):
            with self.subTest(headers=headers, offset=offset):
                self.book([], headers=headers, title_rows=offset)
                with self.assertRaises(WDMSError) as caught:
                    import_sheet(self.path, "任务")
                self.assertEqual(caught.exception.code, "MOYU_SHEET_HEADERS")

    def test_empty_and_over_limit_sheets_are_rejected_before_any_jobs_are_run(self):
        self.book([])
        with self.assertRaises(WDMSError):
            import_sheet(self.path, "任务")
        self.book([(None, "2024-02-29", "000041", None)]*3)
        with patch("moyu_tools.sheet.MAX_TASKS", 2), self.assertRaises(WDMSError):
            import_sheet(self.path, "任务")

    def test_file_changed_during_import_cannot_become_a_resume_identity(self):
        self.book([(None, "2024-02-29", "000041", None)])
        data = self.path.read_bytes()
        with patch("moyu_tools.sheet.source_bytes", side_effect=[data, data+b"changed"]), self.assertRaises(WDMSError) as caught:
            import_sheet(self.path, "任务")
        self.assertEqual(caught.exception.code, "MOYU_SHEET_CHANGED")

    def test_unsupported_and_corrupt_workbooks_have_no_partial_plan(self):
        self.path.write_bytes(b"not a workbook")
        with self.assertRaises(WDMSError):
            import_sheet(self.path, "任务")
        with self.assertRaises(WDMSError):
            import_sheet(self.path.with_suffix(".xls"), "任务")

    def test_identifiers_are_not_rounded_or_reinterpreted(self):
        self.assertEqual(identifier(" 000041 "), "000041")
        self.assertEqual(identifier(41.0), "41")
        self.assertEqual(identifier(41, "000000"), "000041")
        self.assertEqual(identifier("12345678901234567890"), "12345678901234567890")
        for value, fmt in ((True, "General"), (-1, "General"), (1.5, "General"), (float("inf"), "General"),
                           (1234567890123456, "General"), (41, "0.00E+00"), (41, "000-000"), ("41 42", "General")):
            with self.subTest(value=value, fmt=fmt), self.assertRaises(ValueError):
                identifier(value, fmt)

    def test_date_parser_requires_a_full_calendar_date(self):
        self.assertEqual(calendar_date("2024年2月29日"), date(2024, 2, 29))
        for value in ("02/03/24", "2024-02", True, 45351):
            with self.subTest(value=value), self.assertRaises(ValueError):
                calendar_date(value)

    def test_snapshot_roundtrip_does_not_need_the_original_file(self):
        self.book([(None, "2024-02-29", "000041", None)])
        plan = import_sheet(self.path, "任务")
        self.path.unlink()
        self.assertEqual(Plan.parse(plan.document()), plan)
        bad = plan.document()
        bad["rows"].append(dict(bad["rows"][0]))
        with self.assertRaises(WDMSError):
            Plan.parse(bad)


class MoyuSheetFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.engine, self.store, _, _, self.clock, self.output = environment(self.folder, row_plan())
        self.page = Page(self.clock)
        self.remote = NamedRemote()
        self.engine.query, self.engine.remote = self.page.query(), self.remote

    def test_same_day_separate_contracts_unsorted_input_empty_then_next_zip(self):
        self.assertEqual(self.engine.run(trial=True), "trial_passed")
        self.assertEqual(self.engine.run(), "completed")
        self.assertEqual([x["phase"] for x in self.store.value["segments"]], ["done", "empty", "done"])
        self.assertEqual([x["source_row"] for x in self.store.value["segments"]], [2, 3, 5])
        self.assertEqual(self.remote.searches, ["001111", "002222"])
        self.assertEqual([x["contract"] for x in self.page.clicks], ["000041", "000042", "000043"])
        self.assertTrue(all(x["start"] == x["end"] for x in self.page.clicks))
        self.assertEqual([x["start"] for x in self.page.clicks], ["2024-03-02", "2024-02-29", "2024-03-02"])
        exclusions = [p["dates"] for action, p in self.page.actions if action == "conditions"]
        self.assertTrue(exclusions and all(109 in handles for handles in exclusions))
        self.assertEqual(len(list(self.output.glob("*.zip"))), 2)

    def test_empty_first_trial_does_not_unlock_full_run(self):
        self.page.results["000041"] = None
        self.assertEqual(self.engine.run(trial=True), "trial_empty")
        self.assertEqual(self.store.value["segments"][0]["phase"], "empty")
        self.assertFalse(self.store.value["trial_passed"])
        self.assertFalse(self.remote.pastes)
        self.page.results["000042"] = "003333"
        self.assertEqual(self.engine.run(trial=True), "trial_passed")
        self.assertEqual(self.page.clicks[-1]["contract"], "000042")

    def test_wrong_contract_field_and_failed_readback_never_submit_query(self):
        for wrong_label in (True, False):
            with self.subTest(wrong_label=wrong_label):
                self.page.label = "合同号" if wrong_label else "合同编号"
                self.page.changed_readback = not wrong_label
                with self.assertRaises(WDMSError):
                    self.engine.run(trial=True)
                self.assertFalse(self.page.clicks)
                self.assertFalse(self.remote.searches)

    def test_changed_contract_during_query_pauses_before_remote_search(self):
        self.clock.hooks.append(lambda now: self.page.values.update(contract="000099") if now > 3.5 else None)
        with self.assertRaises(WDMSError) as caught:
            self.engine.run(trial=True)
        self.assertEqual(caught.exception.code, "MOYU_CONTRACT_MISMATCH")
        self.assertEqual(len(self.page.clicks), 1)
        self.assertFalse(self.remote.searches)

    def test_fixed_filters_still_pause_before_setting_row_inputs(self):
        self.page.fixed_hash = "c"*64
        with self.assertRaises(WDMSError) as caught:
            self.engine.run(trial=True)
        self.assertEqual(caught.exception.code, "MOYU_CONDITIONS_CHANGED")
        self.assertEqual(self.page.values["contract"], "old")
        self.assertFalse(self.page.clicks)

    def test_visible_column_movement_pauses_instead_of_scrolling(self):
        query = self.page.query()
        self.page.values["contract"] = "000041"
        old = query.record
        query.record = lambda key, context: {"name": "管理编号"} if key == "header" else old(key, context)
        with self.assertRaises(WDMSError) as caught:
            query.snapshot(None, StepContext(20, lambda: False, self.clock))
        self.assertEqual(caught.exception.code, "MOYU_COLUMN_HIDDEN")
        self.assertNotIn("scroll", [action for action, _ in self.page.actions])

    def test_recovery_after_copy_does_not_requery_or_paste_again(self):
        def stop():
            raise WDMSError("CANCELLED", "synthetic stop after paste")
        self.remote.after_paste = stop
        with self.assertRaises(WDMSError):
            self.engine.run(trial=True)
        self.assertEqual(self.store.value["segments"][0]["phase"], "copy_intent")
        self.remote.after_paste = None
        self.assertEqual(self.engine.run(trial=True), "trial_passed")
        self.assertEqual(len(self.page.clicks), 1)
        self.assertEqual(len(self.remote.pastes), 1)

    def test_changed_source_snapshot_or_row_metadata_cannot_reuse_progress(self):
        other = row_plan().document()
        other["rows"][0]["contract"] = "000099"
        with self.assertRaises(WDMSError):
            Store(self.store.folder).load_or_create(Plan.parse(other), self.output, "synthetic-binding")
        value = load_json(self.store.path)
        value["segments"][0]["contract"] = "000099"
        atomic_json_write(self.store.path, value)
        with self.assertRaises(WDMSError):
            Store(self.store.folder).load_or_create(row_plan(), self.output, "synthetic-binding")


class MoyuSheetRuntimeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.env = patch.dict("os.environ", {"LOCALAPPDATA": str(self.folder/"state")})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.output = self.folder/"out"
        self.output.mkdir()
        self.config = {"plan": row_plan().document(), "policy": asdict(Policy(initial_wait=3, settle=1, query_timeout=15)),
                       "output": str(self.output), "task": ""}

    def test_v1_settings_and_progress_are_separate_but_execution_lock_is_shared(self):
        self.assertEqual(state_root(), self.folder/"state"/"MoyuTools"/"sheet_rows_simple")
        self.assertEqual(task_lock_path(), self.folder/"state"/"MoyuTools"/"task.lock")
        self.assertEqual(inventory_lock_path(), self.folder/"state"/"WDMSInventoryDemo"/"workflow.lock")

    def test_spreadsheet_jobs_cannot_use_a_legacy_binding(self):
        with self.assertRaises(WDMSError) as caught:
            make_request(self.config, profile(), trial=True)
        self.assertEqual(caught.exception.code, "MOYU_PLAN_BINDING_MISMATCH")
        self.assertFalse(state_root().exists())

    def test_held_v1_lock_blocks_new_worker_before_native_setup(self):
        path, request = make_request(self.config, sheet_profile(), trial=True)
        with TaskLock(task_lock_path()), patch("moyu_tools.runtime.sys.platform", "win32"), \
                patch("moyu_tools.bridge.Bridge") as native:
            self.assertEqual(worker(path), 2)
        native.assert_not_called()
        self.assertFalse((Path(request["task"])/"progress.json").exists())

    def test_contract_binding_requires_distinct_edit_inside_conditions_and_label_to_its_left(self):
        validate_profile(sheet_profile())
        for change in (lambda p: p["wdms"]["contract"].update(class_name="BUTTON"),
                       lambda p: p["wdms"]["contract"].update(hwnd=p["wdms"]["start"]["hwnd"]),
                       lambda p: p["wdms"]["contract"].update(rect=[700, 280, 800, 300]),
                       lambda p: p["wdms"]["contract_label"].update(point=[850, 195])):
            bad = sheet_profile()
            change(bad)
            with self.assertRaises(WDMSError):
                validate_profile(bad)

    def test_new_worker_route_uses_sheet_snapshot_and_produces_checkable_file_receipt(self):
        path, request = make_request(self.config, sheet_profile(), trial=True)
        clock = Clock()
        page, remote = Page(clock), NamedRemote()
        parent = Mock()
        parent.alive.return_value = True
        def assemble(plan, policy, store, query, files, **kwargs):
            return Engine(plan, policy, store, query, files, clock=clock, sleep=clock.sleep, wall=clock, **kwargs)
        with patch("moyu_tools.runtime.sys.platform", "win32"), patch("moyu_tools.bridge.Bridge"), \
                patch("moyu_tools.bridge.ParentHandle", return_value=parent), patch("moyu_tools.native.Driver"), \
                patch("moyu_tools.native.LocalQuery", return_value=page.query()), \
                patch("moyu_tools.native.RemoteFiles", return_value=remote), patch("moyu_tools.runtime.Engine", side_effect=assemble):
            self.assertEqual(worker(path), 0)
        self.assertEqual(verify_child(0, request), "trial_passed")
        progress = load_json(Path(request["task"])/"progress.json")
        self.assertEqual(progress["segments"][0]["source_row"], 2)
        self.assertEqual(progress["segments"][0]["contract"], "000041")
        self.assertTrue((self.output/"actual_001111.zip").exists())


class MoyuSheetBoundaryTests(unittest.TestCase):
    def test_contract_write_timeout_never_repeats_or_recreates_write_helper(self):
        bridge = Bridge.__new__(Bridge)
        bridge.log, bridge.started = None, True
        bridge.worker = Mock()
        bridge.worker.request.side_effect = WDMSError("WORKER_HARD_TIMEOUT", "synthetic uncertain write")
        with patch("moyu_tools.bridge.JsonLineWorker") as replacement, self.assertRaises(WDMSError) as caught:
            bridge.call("edit", {"value": "000041"})
        self.assertEqual(caught.exception.code, "WORKER_HARD_TIMEOUT")
        replacement.assert_not_called()
        self.assertEqual(bridge.worker.request.call_count, 1)

    def test_contract_read_timeout_only_prepares_a_new_helper_for_future_observation(self):
        bridge = Bridge.__new__(Bridge)
        bridge.log, bridge.started, bridge.close = None, True, Mock()
        old = Mock()
        old.command, old.cwd = ["synthetic"], "."
        old.request.side_effect = WDMSError("WORKER_HARD_TIMEOUT", "synthetic read timeout")
        bridge.worker = old
        with patch("moyu_tools.bridge.JsonLineWorker") as replacement, self.assertRaises(WDMSError) as caught:
            bridge.call("edit", {"hwnd": 109})
        self.assertEqual(caught.exception.code, "MOYU_READ_PENDING")
        replacement.return_value.request.assert_not_called()
        self.assertEqual(old.request.call_count, 1)

    def gui_class(self):
        try:
            from moyu_tools.app import App
        except ModuleNotFoundError as error:
            if error.name in ("tkinter", "_tkinter"):
                self.skipTest("Tk unavailable; GUI state tests only")
            raise
        return App

    def test_selecting_a_different_sheet_requires_import_before_settings_can_start(self):
        App = self.gui_class()
        app = App.__new__(App)
        app.trial, app.run, app.status = Mock(), Mock(), Mock()
        app.config = {"plan": row_plan().document()}
        app.source_changed()
        with self.assertRaises(ValueError):
            app.settings()
        app.run.configure.assert_called_with(state="disabled")
        self.assertEqual(app.config["plan"], row_plan().document())

    def test_failed_import_preserves_prior_plan_and_progress_pointer(self):
        App = self.gui_class()
        app = App.__new__(App)
        app.config = {"plan": row_plan().document(), "task": ""}
        app.messages = queue.Queue()
        app.messages.put(("import", False, "synthetic invalid row"))
        app.importing, app.running, app.quitting = True, False, False
        app.status, app.controls, app.root = Mock(), Mock(), Mock()
        with patch("moyu_tools.app.messagebox.showerror"):
            app.poll()
        self.assertEqual(app.config, {"plan": row_plan().document(), "task": ""})
        self.assertFalse(app.importing)
