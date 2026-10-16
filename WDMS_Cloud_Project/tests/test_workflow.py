import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from desktop_demo import workflow as flow
from errors import NativeTextReadError, WDMSError
from services.task_lock import TaskLock
from native_fixtures import publish_completion, collection_args
from datetime import date


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.folder=Path(self.temp.name)/'flow';self.folder.mkdir()
        self.output=Path(self.temp.name)/'reports';self.output.mkdir()
        self.notice=mock.Mock();self.notice.progress=self.folder/'progress.txt'
        self.order=[]
        self.names=[f'合成工作流货主{i:02d}' for i in range(24)]
        self.owners=self.root/'synthetic_owners.json'
        self.owners.write_text(json.dumps({'owners':self.names},ensure_ascii=False),encoding='utf-8')
        self.config=SimpleNamespace(wdms_exe=str(self.root/'Synthetic.exe'),username='synthetic-user',credential_target='synthetic-target')
        self.store=self.root/'confirmed_owners.json'
        for patch in (mock.patch.object(flow,'load_settings',return_value=self.config),
                      mock.patch.object(flow.owner_discovery,'main',side_effect=self.capture)):
            patch.start();self.addCleanup(patch.stop)

    def capture(self,args):
        args = collection_args(args)
        self.order.append('owners')
        path=Path(args[args.index('--snapshot')+1])
        flow.owner_catalog.write_snapshot(path,self.names,self.config,{'pid':7,'hwnd':100})
        folder=Path(args[args.index('--out')+1])
        flow.probe.record(folder/'stage.jsonl','owner_list_capture_complete',count=24)
        return 0

    def stage(self,name,marker):
        def run(args):
            args = collection_args(args)
            self.order.append(name)
            folder=Path(args[args.index('--out')+1])
            flow.probe.record(folder/'stage.jsonl',marker)
            return 0
        return run

    def collection(self,args):
        args = collection_args(args)
        self.order.append('collection')
        self.assertIn('--managed-notices',args)
        owners=Path(args[args.index('--owners')+1])
        self.assertEqual(json.loads(owners.read_text(encoding='utf-8'))['owners'],self.names)
        path=self.output/'native_trial_fixture'/'result.xlsx'
        count=int(args[args.index('--limit')+1])
        flow.trial.write_trial_excel([{'owner':name,'whole':'1','loose':'2.54'} for name in self.names[:count]],path, task_date=date.fromisoformat(args[args.index('--report-date')+1]))
        manifest=Path(args[args.index('--result-manifest')+1])
        manifest.write_text(json.dumps({'report':str(path),'row_count':count,'requires_manual_review':True,'report_date':args[args.index('--report-date')+1],'run_id':args[args.index('--run-id')+1]}),encoding='utf-8')
        return 0

    def navigation_with_page_timeout(self,recover):
        root={'hwnd':100,'pid':7,'class_name':'WindowsForms10.Window.8.synthetic',
              'visible':True,'minimized':False,'rect':[0,0,1000,800]}
        page=[]
        for hwnd,name,kind,rect in (
            (10,'在库商品检索','STATIC',[20,50,120,70]),
            (11,'货主','STATIC',[0,100,40,120]),
            (12,'','COMBOBOX',[45,100,160,120]),
            (13,'查询(&F)','BUTTON',[200,100,260,120]),
            (14,'整件数：','STATIC',[0,200,50,220]),
            (15,'散件数：','STATIC',[100,200,150,220])):
            node={'hwnd':hwnd,'pid':7,'class_name':f'WindowsForms10.{kind}.synthetic',
                  'visible':True,'rect':rect,'event':'native_control'}
            page.extend([node,{**node,'event':'native_text','text':name}])
        nodes={e['hwnd']:e for e in page if e['event']=='native_control'};nodes[100]=root
        names={e['hwnd']:e['text'] for e in page if e['event']=='native_text'}
        self.navigation_api=SimpleNamespace(metadata=lambda hwnd:nodes[hwnd],
            read_label_text=lambda data,log:names[data['hwnd']],health=mock.Mock(),
            client_screen_rect=lambda hwnd:nodes[hwnd]['rect'],
            user32=SimpleNamespace(IsWindowEnabled=lambda hwnd:True,
                                   GetParent=lambda hwnd:0 if hwnd==100 else 100),
            click_navigation_input=mock.Mock(side_effect=lambda root,target,log,before_press:before_press()))
        error=NativeTextReadError('NATIVE_TEXT_READ_FAILED','合成页面读取超时',winerror=1460,
                                  hwnd=90,class_name='WindowsForms10.STATIC.synthetic')
        def run(args):
            args = collection_args(args)
            self.order.append('navigation')
            log=Path(args[args.index('--out')+1])/'stage.jsonl'
            after=[0]
            def scan(stage):
                if stage=='initial':return page[:2]
                after[0]+=1
                if not recover or after[0]==1:raise error
                return page
            try:
                flow.navigation.navigate(self.navigation_api,root,log,scan)
                flow.probe.record(log,'automatic_navigation_complete',query_performed=False)
                return 0
            except WDMSError as failure:
                flow.probe.record(log,'probe_error',**flow.probe.error_info(failure))
                return 2
        return run

    def test_navigation_read_recovery_continues_to_owners_and_real_excel(self):
        clock=[0.0]
        with mock.patch.object(flow.navigation.time,'monotonic',side_effect=lambda:clock[0]), \
             mock.patch.object(flow.navigation.time,'sleep',side_effect=lambda s:clock.__setitem__(0,clock[0]+s)), \
             mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')), \
             mock.patch.object(flow.navigation,'main',side_effect=self.navigation_with_page_timeout(True)), \
             mock.patch.object(flow.trial,'run_collection',side_effect=self.collection):
            report=flow.run_stages(self.folder,self.output,3,self.notice)
        self.assertEqual(self.order,['login','navigation','owners','collection'])
        self.assertTrue(report.is_file());self.assertFalse(self.store.exists())
        self.notice.close.assert_not_called();self.notice.restart.assert_not_called()
        self.assertEqual(self.navigation_api.click_navigation_input.call_count,1)
        nav_events=flow.stage_events(self.folder/'navigation')
        self.assertTrue(any(e['event']=='page_scan_recovered' for e in nav_events))
        flow.verify_report(report,tuple(self.names[:3]))

    def test_navigation_read_timeout_cannot_reach_owner_capture_or_query(self):
        clock=[0.0]
        with mock.patch.object(flow.navigation.time,'monotonic',side_effect=lambda:clock[0]), \
             mock.patch.object(flow.navigation.time,'sleep',side_effect=lambda s:clock.__setitem__(0,clock[0]+s)), \
             mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')), \
             mock.patch.object(flow.navigation,'main',side_effect=self.navigation_with_page_timeout(False)), \
             mock.patch.object(flow.trial,'run_collection') as collection:
            with self.assertRaises(WDMSError) as caught:
                flow.run_stages(self.folder,self.output,3,self.notice)
        self.assertEqual(caught.exception.code,'NATIVE_TEXT_READ_FAILED')
        self.assertEqual(self.order,['login','navigation'])
        collection.assert_not_called()
        self.assertFalse((self.folder/'owners_snapshot.json').exists())
        self.assertEqual(self.navigation_api.click_navigation_input.call_count,1)

    def test_success_requires_all_stages_and_the_report_manifest(self):
        with mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')),mock.patch.object(flow.navigation,'main',side_effect=self.stage('navigation','automatic_navigation_complete')),mock.patch.object(flow.trial,'run_collection',side_effect=self.collection):
            report=flow.run_stages(self.folder,self.output,24,self.notice)
        self.assertEqual(self.order,['login','navigation','owners','collection'])
        self.assertTrue(report.is_file())
        self.assertEqual(flow.probe.read_events(self.folder/'flow.jsonl')[-1]['event'],'automatic_workflow_complete')

    def test_recovered_read_warning_is_not_used_as_terminal_stage_failure(self):
        warning={'event':'query_settle_read_unavailable','error_code':'NATIVE_TEXT_READ_FAILED','error':'synthetic recovered read'}
        error=flow.stage_failure([warning],'WORKFLOW_COLLECTION_FAILED','synthetic missing result')
        self.assertEqual(error.code,'WORKFLOW_COLLECTION_FAILED')
        terminal={'event':'trial_worker_failed','error_code':'TRIAL_WORKER_EXITED','error':'synthetic exit 7'}
        error=flow.stage_failure([terminal,warning],'WORKFLOW_COLLECTION_FAILED','fallback')
        self.assertEqual((error.code,error.message),('TRIAL_WORKER_EXITED','synthetic exit 7'))

    def test_collection_logs_are_copied_without_snapshots_reports_or_old_runs(self):
        old=self.output/'native_trial_old';old.mkdir();(old/'owner_01.jsonl').write_text('old evidence', encoding='utf-8')
        def collection(args):
            code=self.collection(args);folder=self.output/'native_trial_fixture'
            for name in ('owner_01.jsonl','owner_01.stderr.txt','trial_error.jsonl'):
                (folder/name).write_text('synthetic evidence', encoding='utf-8')
            for name in ('owners_snapshot.json','trial_results.json','settings.json'):
                (folder/name).write_text('must not copy', encoding='utf-8')
            return code
        self.run_flow(collection=collection,limit=3)
        copied=self.folder/'collection'
        self.assertEqual({p.name for p in copied.iterdir()},{'owner_01.jsonl','owner_01.stderr.txt','trial_error.jsonl'})
        self.assertEqual((copied/'owner_01.jsonl').read_text(encoding='utf-8'),'synthetic evidence')

    def test_copy_failure_preserves_the_original_collection_error(self):
        def collect(args):
            folder=self.output/'native_trial_failed';folder.mkdir()
            flow.probe.record(folder/'owner_01.jsonl','probe_error',error_code='QUERY_OWNER_MISMATCH',error='synthetic mismatch')
            return 2
        with mock.patch.object(flow.shutil,'copyfile',side_effect=OSError('synthetic copy failure')):
            with self.assertRaises(WDMSError) as caught:self.run_flow(collection=collect)
        self.assertEqual(caught.exception.code,'QUERY_OWNER_MISMATCH')
        self.assertTrue(any(e['event']=='collection_evidence_copy_failed' for e in flow.probe.read_events(self.folder/'flow.jsonl')))

    def test_three_completed_owners_then_unlogged_fourth_exit_has_correct_workflow_error(self):
        import contextlib,io
        real_collector=flow.trial.run_collection
        def worker(command,log,timeout,cancelled=None,stderr_path=None):
            index=int(command[command.index('--index')+1])
            stderr_path.write_text('synthetic fatal worker trace' if index==3 else '', encoding='utf-8')
            if index==3:return 7
            if index==2:
                flow.probe.record(log,'query_settle_read_unavailable',error_code='NATIVE_TEXT_READ_FAILED',error='synthetic recovered text read')
            flow.probe.record(log,'trial_owner_complete',owner=self.names[index],whole='1',loose='2.54',total='3.54')
            publish_completion(command,self.names[index])
            return 0
        with mock.patch.object(flow.trial.probe,'supervise',side_effect=worker) as supervisor, \
             mock.patch.object(flow.trial,'show_notice'),contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(WDMSError) as caught:self.run_flow(collection=real_collector)
        self.assertEqual(caught.exception.code,'TRIAL_WORKER_EXITED')
        self.assertIn('第4个货主',caught.exception.message);self.assertIn('退出码为7',caught.exception.message)
        self.assertEqual(supervisor.call_count,4)
        copied=self.folder/'collection'
        self.assertEqual((copied/'owner_04.stderr.txt').read_text(encoding='utf-8'),'synthetic fatal worker trace')
        self.assertEqual(flow.probe.read_events(copied/'owner_04.jsonl')[-1]['error_code'],'TRIAL_WORKER_EXITED')
        self.assertFalse(list(self.output.rglob('*.xlsx')));self.assert_no_success()

    def test_login_failure_stops_navigation_and_collection(self):
        with mock.patch.object(flow.login,'main',return_value=2),mock.patch.object(flow.navigation,'main') as nav,mock.patch.object(flow.trial,'run_collection') as trial:
            with self.assertRaises(WDMSError):flow.run_stages(self.folder,self.output,24,self.notice)
        nav.assert_not_called();trial.assert_not_called()

    def test_zero_exit_without_login_marker_cannot_start_navigation(self):
        with mock.patch.object(flow.login,'main',return_value=0),mock.patch.object(flow.navigation,'main') as nav,mock.patch.object(flow.trial,'run_collection') as trial:
            with self.assertRaises(WDMSError) as caught:flow.run_stages(self.folder,self.output,24,self.notice)
        self.assertEqual(caught.exception.code,'WORKFLOW_STAGE_UNCONFIRMED')
        nav.assert_not_called();trial.assert_not_called()

    def test_navigation_failure_cannot_start_collection(self):
        with mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')),mock.patch.object(flow.navigation,'main',return_value=2),mock.patch.object(flow.trial,'run_collection') as trial:
            with self.assertRaises(WDMSError):flow.run_stages(self.folder,self.output,24,self.notice)
        trial.assert_not_called()

    def test_report_missing_prevents_success_marker(self):
        def collection(args):
            args = collection_args(args)
            path=Path(args[args.index('--result-manifest')+1])
            path.write_text(json.dumps({'report':str(self.output/'missing.xlsx'),'row_count':24}), encoding='utf-8')
            return 0
        with mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')),mock.patch.object(flow.navigation,'main',side_effect=self.stage('navigation','automatic_navigation_complete')),mock.patch.object(flow.trial,'run_collection',side_effect=collection):
            with self.assertRaises(WDMSError):flow.run_stages(self.folder,self.output,24,self.notice)
        self.assertFalse(any(e['event']=='automatic_workflow_complete' for e in flow.probe.read_events(self.folder/'flow.jsonl')))

    def test_deadline_and_cancel_path_are_passed_to_all_stages(self):
        cancel_file=self.folder/'cancel.request'
        captured=[]

        def stage(name,marker):
            def run(args):
                args = collection_args(args)
                captured.append(args)
                folder=Path(args[args.index('--out')+1])
                flow.probe.record(folder/'stage.jsonl',marker)
                return 0
            return run

        with mock.patch.object(flow.login,'main',side_effect=stage('login','automatic_login_complete')), \
             mock.patch.object(flow.navigation,'main',side_effect=stage('navigation','automatic_navigation_complete')), \
             mock.patch.object(flow.owner_discovery,'main',side_effect=lambda args:captured.append(args) or self.capture(args)), \
             mock.patch.object(flow.trial,'run_collection',side_effect=lambda args: captured.append(args) or self.collection(args)):
            flow.run_stages(self.folder,self.output,24,self.notice,cancel_file=cancel_file)

        self.assertEqual(len(captured),4)
        for args in captured:
            args = collection_args(args)
            self.assertEqual(args[args.index('--cancel-file')+1],str(cancel_file.resolve()))
            self.assertGreater(float(args[args.index('--execution-deadline')+1]),0)

    def test_existing_cancel_file_prevents_login(self):
        cancel_file=self.folder/'cancel.request'
        cancel_file.touch()
        with mock.patch.object(flow.login,'main') as login, mock.patch.object(flow.navigation,'main') as nav, \
             mock.patch.object(flow.trial,'run_collection') as trial:
            with self.assertRaises(WDMSError) as caught:
                flow.run_stages(self.folder,self.output,24,self.notice,cancel_file=cancel_file)
        self.assertEqual(caught.exception.code,'TASK_CANCELLED')
        login.assert_not_called();nav.assert_not_called();trial.assert_not_called()
        self.assertTrue(cancel_file.exists())

    def test_cancel_created_after_login_prevents_navigation(self):
        cancel_file=self.folder/'cancel.request'

        def login_stage(args):
            args = collection_args(args)
            directory=Path(args[args.index('--out')+1])
            flow.probe.record(directory/'stage.jsonl','automatic_login_complete')
            cancel_file.touch()
            return 0

        with mock.patch.object(flow.login,'main',side_effect=login_stage), \
             mock.patch.object(flow.navigation,'main') as nav, mock.patch.object(flow.trial,'run_collection') as trial:
            with self.assertRaises(WDMSError) as caught:
                flow.run_stages(self.folder,self.output,24,self.notice,cancel_file=cancel_file)
        self.assertEqual(caught.exception.code,'TASK_CANCELLED')
        nav.assert_not_called();trial.assert_not_called()

    def test_missing_corrupt_and_invalid_manifest_never_mark_workflow_complete(self):
        cases=(('missing',None),('corrupt','{broken json'),
               ('wrong-fields',json.dumps({'report':'','row_count':True})))
        for name,content in cases:
            with self.subTest(manifest=name):
                folder=self.root/name
                folder.mkdir()

                def collection(args):
                    args = collection_args(args)
                    manifest=Path(args[args.index('--result-manifest')+1])
                    if content is not None:
                        manifest.write_text(content,encoding='utf-8')
                    return 0

                with mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')), \
                     mock.patch.object(flow.navigation,'main',side_effect=self.stage('navigation','automatic_navigation_complete')), \
                     mock.patch.object(flow.trial,'run_collection',side_effect=collection):
                    with self.assertRaises(WDMSError) as caught:
                        flow.run_stages(folder,self.output,24,self.notice)
                self.assertEqual(caught.exception.code,'WORKFLOW_REPORT_UNCONFIRMED')
                events=flow.probe.read_events(folder/'flow.jsonl')
                self.assertFalse(any(event['event']=='automatic_workflow_complete' for event in events))

    def test_main_lock_occupation_returns_two_without_login(self):
        settings=self.root/'settings.json'
        output=self.root/'configured-output'
        workflows=self.root/'workflows'
        occupied=TaskLock(settings.parent/'workflow.lock')
        occupied.acquire()
        try:
            with mock.patch.object(flow,'settings_path',return_value=settings), \
                 mock.patch.object(flow,'load_settings') as load_settings, \
                 mock.patch.object(flow,'validate_output',return_value=output), \
                 mock.patch.object(flow.trial,'show_notice') as show_notice, \
                 mock.patch.object(flow.login,'main') as login:
                code=flow.main(['--out',str(workflows)])
        finally:
            occupied.release()
        self.assertEqual(code,2)
        login.assert_not_called();load_settings.assert_not_called();show_notice.assert_called_once()
        flow_dirs=list(workflows.glob('flow_*'))
        self.assertEqual(len(flow_dirs),1)
        errors=flow.probe.read_events(flow_dirs[0]/'flow.jsonl')
        self.assertEqual(errors[-1]['error_code'],'TASK_ALREADY_RUNNING')

    def test_notice_closes_and_lock_releases_after_exception_and_interrupt(self):
        settings=self.root/'settings.json'
        output=self.root/'configured-output'
        for failure in (RuntimeError('test failure'),KeyboardInterrupt()):
            with self.subTest(failure=type(failure).__name__):
                notice=mock.Mock()
                workflows=self.root/('flows-'+type(failure).__name__)
                cancel_file=self.root/(type(failure).__name__+'.cancel')
                with mock.patch.object(flow,'settings_path',return_value=settings), \
                     mock.patch.object(flow,'load_settings',return_value=SimpleNamespace(output_dir=output)), \
                     mock.patch.object(flow,'validate_output',return_value=output), \
                     mock.patch.object(flow,'preflight',return_value=flow.WorkflowInputs(output,tuple(self.names))), \
                     mock.patch.object(flow.trial,'RunningNotice',return_value=notice), \
                     mock.patch.object(flow.trial,'show_notice'), \
                     mock.patch.object(flow,'run_stages',side_effect=failure):
                    code=flow.main(['--out',str(workflows),'--cancel-file',str(cancel_file)])
                self.assertEqual(code,2)
                notice.start.assert_called_once_with();notice.close.assert_called_once_with()
                self.assertFalse(TaskLock.is_locked(settings.parent/'workflow.lock'))
                flow_dir=next(workflows.glob('flow_*'))
                start=flow.probe.read_events(flow_dir/'flow.jsonl')[0]
                self.assertEqual(start['revision'],flow.WORKFLOW_REVISION)
                self.assertEqual(start['folder'],flow_dir.name)
                self.assertEqual(start['cancel_file'],str(cancel_file.resolve()))

    def run_flow(self, collection=None, limit=24, login=None):
        with mock.patch.object(flow.login,'main',side_effect=login or self.stage('login','automatic_login_complete')), \
             mock.patch.object(flow.navigation,'main',side_effect=self.stage('navigation','automatic_navigation_complete')), \
             mock.patch.object(flow.trial,'run_collection',side_effect=collection or self.collection):
            return flow.run_stages(self.folder,self.output,limit,self.notice)

    def assert_no_success(self):
        self.assertFalse(any(event['event']=='automatic_workflow_complete'
                             for event in flow.probe.read_events(self.folder/'flow.jsonl')))

    def test_three_owner_flow_uses_first_three_snapshot_names(self):
        report=self.run_flow(limit=3)
        from openpyxl import load_workbook
        book=load_workbook(report,read_only=True)
        try:
            self.assertEqual([row[0] for row in book.active.iter_rows(min_row=2,values_only=True)],self.names[:3])
        finally:
            book.close()

    def test_workflow_verifies_the_start_date_after_midnight(self):
        from services import trial_excel_reporter
        with mock.patch.object(flow,'date') as clock_date, mock.patch.object(trial_excel_reporter,'date') as writer_date:
            clock_date.today.side_effect=[date(2026,10,13),date(2026,10,14)]
            writer_date.today.return_value=date(2026,10,14)
            report=self.run_flow(limit=3)
        trial_excel_reporter.verify_report(report,self.names[:3],date(2026,10,13))

    def test_source_json_changed_during_login_is_ignored_by_automatic_read(self):
        def login(args):
            self.owners.write_text(json.dumps({'owners':[f'另一轮合成货主{i}' for i in range(24)]}),encoding='utf-8')
            return self.stage('login','automatic_login_complete')(args)
        report=self.run_flow(login=login)
        self.assertTrue(report.is_file())
        self.assertEqual(json.loads((self.folder/'owners_snapshot.json').read_text(encoding='utf-8'))['owners'],self.names)

    def test_first_and_repeated_collection_run_without_owner_dialog_pause(self):
        self.run_flow()
        self.folder=self.root/'flow_second';self.folder.mkdir()
        self.output=self.root/'reports_second';self.output.mkdir()
        self.run_flow()
        self.assertFalse(self.store.exists())
        self.notice.close.assert_not_called();self.notice.restart.assert_not_called()

    def test_cancelling_after_capture_still_never_starts_collection(self):
        def capture(args):
            result=self.capture(args)
            (self.folder/'cancel.request').touch()
            return result
        with mock.patch.object(flow.owner_discovery,'main',side_effect=capture):
            with self.assertRaises(WDMSError) as caught:self.run_flow()
        self.assertEqual(caught.exception.code,'TASK_CANCELLED')
        self.assertEqual(self.order,['login','navigation','owners']);self.assert_no_success()
        self.assertFalse(self.store.exists());self.notice.restart.assert_not_called()

    def test_changed_list_runs_automatically_and_excel_contains_new_names(self):
        self.run_flow()
        self.store.write_text('{old confirmation is irrelevant}', encoding='utf-8')
        previous=self.store.read_bytes()
        self.names[-1]='变化后的合成货主'
        self.names.reverse()
        self.folder=self.root/'flow_changed';self.folder.mkdir()
        self.output=self.root/'reports_changed';self.output.mkdir()
        self.order.clear()
        report=self.run_flow()
        self.assertEqual(self.order,['login','navigation','owners','collection'])
        self.assertEqual(self.store.read_bytes(),previous)
        flow.verify_report(report,tuple(self.names))
        self.notice.close.assert_not_called();self.notice.restart.assert_not_called()

    def test_missing_capture_marker_never_collects(self):
        with mock.patch.object(flow.owner_discovery,'main',return_value=0):
            with self.assertRaises(WDMSError) as caught:self.run_flow()
        self.assertEqual(caught.exception.code,'WORKFLOW_STAGE_UNCONFIRMED')
        self.assert_no_success()

    def test_changed_config_after_capture_does_not_start_collection(self):
        def capture(args):
            result=self.capture(args)
            self.config.username='changed-after-capture'
            return result
        with mock.patch.object(flow.owner_discovery,'main',side_effect=capture):
            with self.assertRaises(WDMSError) as caught:self.run_flow()
        self.assertEqual(caught.exception.code,'OWNER_SESSION_CHANGED')
        self.assertEqual(self.order,['login','navigation','owners']);self.assert_no_success()

    def test_existing_stage_artifacts_prevent_reusing_old_results(self):
        old=self.folder/'collection_result.json';old.write_text('old result', encoding='utf-8')
        with mock.patch.object(flow.login,'main') as login:
            with self.assertRaises(WDMSError) as caught:
                flow.run_stages(self.folder,self.output,24,self.notice)
        self.assertEqual(caught.exception.code,'WORKFLOW_ARTIFACTS_EXIST')
        self.assertEqual(old.read_text(encoding='utf-8'),'old result');login.assert_not_called()

    def test_zero_exit_with_error_and_marker_preserves_original_error(self):
        def login(args):
            args = collection_args(args)
            folder=Path(args[args.index('--out')+1])
            flow.probe.record(folder/'stage.jsonl','automatic_login_complete')
            flow.probe.record(folder/'stage.jsonl','probe_error',error_code='LOGIN_CREDENTIAL_MISMATCH',error='合成凭据不一致')
            return 0
        with mock.patch.object(flow.navigation,'main') as navigation:
            with self.assertRaises(WDMSError) as caught:
                self.run_flow(login=login)
        self.assertEqual(caught.exception.code,'LOGIN_CREDENTIAL_MISMATCH')
        self.assert_no_success()

    def test_duplicate_success_markers_are_not_accepted(self):
        def login(args):
            run=self.stage('login','automatic_login_complete')
            run(args);return run(args)
        with self.assertRaises(WDMSError) as caught:self.run_flow(login=login)
        self.assertEqual(caught.exception.code,'WORKFLOW_STAGE_UNCONFIRMED')
        self.assertEqual(self.order,['login','login']);self.assert_no_success()

    def test_recovered_health_warning_does_not_block_confirmed_success(self):
        def login(args):
            args = collection_args(args)
            folder=Path(args[args.index('--out')+1])
            flow.probe.record(folder/'stage.jsonl','health_wait',error_code='WDMS_HEALTH_CHECK_FAILED')
            flow.probe.record(folder/'stage.jsonl','health_recovered')
            return self.stage('login','automatic_login_complete')(args)
        self.assertTrue(self.run_flow(login=login).is_file())

    def test_auxiliary_diagnostic_error_does_not_override_confirmed_navigation(self):
        def navigation(args):
            args = collection_args(args)
            folder=Path(args[args.index('--out')+1])
            flow.probe.record(folder/'stage.jsonl','navigation_menu_diagnostic_error',error_code='MENU_DIAGNOSTIC_PARTIAL')
            return self.stage('navigation','automatic_navigation_complete')(args)
        with mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')), \
             mock.patch.object(flow.navigation,'main',side_effect=navigation), \
             mock.patch.object(flow.trial,'run_collection',side_effect=self.collection):
            self.assertTrue(flow.run_stages(self.folder,self.output,24,self.notice).is_file())

    def test_invalid_limit_stops_before_any_stage(self):
        for limit in (True,3.0,None,0,25):
            with self.subTest(limit=limit),mock.patch.object(flow.login,'main') as login:
                with self.assertRaises(WDMSError) as caught:
                    flow.run_stages(self.folder,self.output,limit,self.notice)
                self.assertEqual(caught.exception.code,'WORKFLOW_LIMIT_INVALID');login.assert_not_called()

    def test_system_exit_is_a_recorded_stage_failure(self):
        def login(args):raise SystemExit(2)
        with self.assertRaises(WDMSError) as caught:self.run_flow(login=login)
        self.assertEqual(caught.exception.code,'WORKFLOW_STAGE_EXITED')
        self.assertEqual(flow.probe.read_events(self.folder/'flow.jsonl')[-1]['event'],'step_failed')
        self.assert_no_success()

    def test_collection_failure_preserves_this_runs_worker_error(self):
        def collection(args):
            folder=self.output/'native_trial_failed';folder.mkdir()
            flow.probe.record(folder/'owner_01.jsonl','probe_error',error_code='OWNER_SELECTION_MISMATCH',error='合成货主选择不一致')
            return 2
        with self.assertRaises(WDMSError) as caught:self.run_flow(collection)
        self.assertEqual(caught.exception.code,'OWNER_SELECTION_MISMATCH')
        self.assert_no_success()

    def test_corrupt_excel_is_rejected_despite_valid_manifest(self):
        def collection(args):
            self.collection(args)
            (self.output/'native_trial_fixture'/'result.xlsx').write_bytes(b'not an Excel file')
            return 0
        with self.assertRaises(WDMSError) as caught:self.run_flow(collection)
        self.assertEqual(caught.exception.code,'WORKFLOW_REPORT_UNCONFIRMED');self.assert_no_success()

    def test_old_report_cannot_be_substituted_for_this_run(self):
        old=self.output/'native_trial_old'/'result.xlsx'
        flow.trial.write_trial_excel([{'owner':name,'whole':9,'loose':8} for name in self.names],old)
        old_bytes=old.read_bytes()
        def collection(args):
            args = collection_args(args)
            self.collection(args)
            manifest=Path(args[args.index('--result-manifest')+1])
            manifest.write_text(json.dumps({'report':str(old),'row_count':24,'requires_manual_review':True}), encoding='utf-8')
            return 0
        with self.assertRaises(WDMSError) as caught:self.run_flow(collection)
        self.assertEqual(caught.exception.code,'WORKFLOW_REPORT_UNCONFIRMED')
        self.assertEqual(old.read_bytes(),old_bytes);self.assert_no_success()

    def test_wrong_owner_order_and_partial_workbook_are_rejected(self):
        for name,rows in (('reversed',list(reversed(self.names))),('partial',self.names[:3])):
            with self.subTest(case=name):
                self.folder=self.root/('flow-'+name);self.folder.mkdir()
                self.output=self.root/('reports-'+name);self.output.mkdir()
                def collection(args):
                    self.collection(args)
                    flow.trial.write_trial_excel([{'owner':owner,'whole':1,'loose':2} for owner in rows],
                                                self.output/'native_trial_fixture'/'result.xlsx')
                    return 0
                with self.assertRaises(WDMSError) as caught:self.run_flow(collection)
                self.assertEqual(caught.exception.code,'WORKFLOW_REPORT_UNCONFIRMED');self.assert_no_success()

    def test_run_stages_with_real_collector_and_synthetic_worker(self):
        import contextlib,io
        def worker(command,log,timeout,cancelled=None,stderr_path=None):
            names=json.loads(Path(command[command.index('--owners')+1]).read_text(encoding='utf-8'))['owners']
            index=int(command[command.index('--index')+1])
            flow.probe.record(log,'trial_owner_complete',owner=names[index],whole='1',loose='2.54')
            publish_completion(command,names[index])
            return 0
        with mock.patch.object(flow.login,'main',side_effect=self.stage('login','automatic_login_complete')), \
             mock.patch.object(flow.navigation,'main',side_effect=self.stage('navigation','automatic_navigation_complete')), \
             mock.patch.object(flow.probe,'supervise',side_effect=worker), \
             mock.patch.object(flow.trial,'show_notice') as popup,contextlib.redirect_stdout(io.StringIO()):
            report=flow.run_stages(self.folder,self.output,3,self.notice)
        self.assertTrue(report.is_file());popup.assert_not_called()
        self.assertTrue(json.loads((self.folder/'collection_result.json').read_text(encoding='utf-8'))['requires_manual_review'])


class WorkflowPreflightTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.exe=self.root/'SyntheticWDMS.exe';self.exe.write_bytes(b'synthetic fixture only')
        self.owners=self.root/'synthetic_owners.json'
        self.names=[f'合成预检货主{i}' for i in range(24)]
        self.owners.write_text(json.dumps({'owners':self.names}), encoding='utf-8')
        self.config=SimpleNamespace(wdms_exe=str(self.exe),output_dir=str(self.root/'reports'),username='synthetic-user',credential_target='synthetic-target')
        self.notice=mock.Mock()

    def main(self, extra=()):
        with mock.patch.object(flow,'DEFAULT_OWNERS_PATH',self.owners), \
             mock.patch.object(flow,'settings_path',return_value=self.root/'settings.json'), \
             mock.patch.object(flow,'load_settings',return_value=self.config), \
             mock.patch.object(flow.probe,'inspect_environment',return_value={'integrity_rid':8192}), \
             mock.patch.object(flow.trial,'RunningNotice',return_value=self.notice) as banner, \
             mock.patch.object(flow.trial,'show_notice') as popup, \
             mock.patch.object(flow.login,'main') as login, \
             mock.patch.object(flow.navigation,'main') as navigation, \
             mock.patch.object(flow.trial,'run_collection') as collection:
            code=flow.main(['--out',str(self.root/'flows'),'--limit','3',*extra])
        folder=next((self.root/'flows').glob('flow_*'))
        return code,flow.probe.read_events(folder/'flow.jsonl'),login,navigation,collection,banner,popup

    def test_missing_source_json_does_not_block_automatic_workflow(self):
        self.owners.unlink()
        code,events,login,navigation,collection,banner,popup=self.main()
        self.assertEqual(code,2)  # This fixture's login runner is deliberately unconfirmed.
        self.assertEqual(events[-1]['error_code'],'WORKFLOW_STAGE_FAILED')
        self.assertTrue(any(event['event']=='workflow_preflight_complete' for event in events))
        login.assert_called_once();navigation.assert_not_called();collection.assert_not_called();banner.assert_called_once()

    def test_explicit_missing_list_still_fails_before_wdms_interaction(self):
        self.owners.unlink()
        code,events,login,navigation,collection,banner,popup=self.main(('--owners',str(self.owners)))
        self.assertEqual(code,2);self.assertEqual(events[-1]['error_code'],'WORKFLOW_OWNERS_MISSING')
        login.assert_not_called();banner.assert_not_called()

    def test_missing_executable_stops_before_interacting_with_wdms(self):
        self.exe.unlink()
        code,events,login,navigation,collection,banner,popup=self.main()
        self.assertEqual(code,2);self.assertEqual(events[-1]['error_code'],'WORKFLOW_EXECUTABLE_INVALID')
        login.assert_not_called();banner.assert_not_called()

    def test_unwritable_output_stops_before_interacting_with_wdms(self):
        with mock.patch.object(flow,'validate_output',side_effect=ValueError('private detail')):
            code,events,login,navigation,collection,banner,popup=self.main()
        self.assertEqual(code,2);self.assertEqual(events[-1]['error_code'],'WORKFLOW_OUTPUT_INVALID')
        login.assert_not_called();banner.assert_not_called()
        self.assertNotIn('private detail',json.dumps(events));self.assertFalse(TaskLock.is_locked(self.root/'workflow.lock'))

    def test_malformed_owner_lists_are_rejected_and_bom_is_accepted(self):
        cases=(None,'text',[],self.names[:23],self.names[:23]+[self.names[0]],self.names[:23]+[None],self.names[:23]+[' '])
        for names in cases:
            with self.subTest(names_type=type(names).__name__):
                self.owners.write_text(json.dumps({'owners':names}), encoding='utf-8')
                with self.assertRaises(WDMSError) as caught:flow.read_owners(self.owners)
                self.assertEqual(caught.exception.code,'WORKFLOW_OWNERS_INVALID')
        self.owners.write_text(json.dumps({'owners':self.names}),encoding='utf-8-sig')
        self.assertEqual(flow.read_owners(self.owners),tuple(self.names))

    def test_notice_cleanup_failure_does_not_replace_original_stage_error(self):
        self.notice.close.side_effect=RuntimeError('private cleanup detail')
        inputs=flow.WorkflowInputs(self.root/'reports',tuple(self.names))
        with mock.patch.object(flow,'preflight',return_value=inputs), \
             mock.patch.object(flow,'run_stages',side_effect=WDMSError('OWNER_SELECTION_MISMATCH','合成选择错误')):
            code,events,login,navigation,collection,banner,popup=self.main()
        self.assertEqual(code,2)
        errors=[event['error_code'] for event in events if event.get('error_code')]
        self.assertEqual(errors,['OWNER_SELECTION_MISMATCH','WORKFLOW_NOTICE_CLEANUP_FAILED'])
        self.assertIn('合成选择错误',popup.call_args.args[1])
        self.assertNotIn('private cleanup detail',popup.call_args.args[1])
        self.assertFalse(TaskLock.is_locked(self.root/'workflow.lock'))

if __name__=='__main__':unittest.main()
