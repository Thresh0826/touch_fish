import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from scripts import run_native_trial as trial
from errors import WDMSError
from services import trial_excel_reporter
from services import owner_catalog
from native_fixtures import isolate_trial_notices, publish_completion, publish_failure

class NativeTrialTests(unittest.TestCase):
    def setUp(self):
        self.notice_factory, self.notice_popup = isolate_trial_notices(self)
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.log=Path(self.temp.name)/'test.jsonl'
        self.owner_names=[f'合成测试货主{i:02d}' for i in range(24)]
        self.owners=Path(self.temp.name)/'synthetic_owners.json'
        self.owners.write_text(json.dumps({'owners':self.owner_names},ensure_ascii=False),encoding='utf-8')

    def successful_worker(self,command,log,timeout,cancelled=None,stderr_path=None):
        index=int(command[command.index('--index')+1])
        trial.probe.record(log,'trial_owner_complete',owner=self.owner_names[index],whole='1',loose='2.54',total='3.54')
        publish_completion(command,self.owner_names[index])
        return 0

    def automatic_capture(self,args):
        owner_catalog.write_snapshot(Path(args[args.index('--snapshot')+1]),self.owner_names,
                                     self.auto_config,{'pid':7,'hwnd':100})
        return 0

    def test_default_collection_reads_names_without_source_json_or_owner_dialog(self):
        import contextlib,io
        from desktop_demo import owner_discovery,settings
        root=Path(self.temp.name)
        self.auto_config=SimpleNamespace(wdms_exe=str(root/'Synthetic.exe'),username='synthetic-user',credential_target='synthetic-target')
        def worker(command,log,timeout,cancelled=None,stderr_path=None):
            owner_catalog.read_plan(command[command.index('--owners')+1],command[command.index('--owners-sha256')+1])
            return self.successful_worker(command,log,timeout,cancelled)
        with mock.patch.object(settings,'load_settings',return_value=self.auto_config), \
             mock.patch.object(owner_discovery,'main',side_effect=self.automatic_capture) as capture, \
             mock.patch.object(trial.probe,'supervise',side_effect=worker), \
             mock.patch.object(trial,'show_notice'),contextlib.redirect_stdout(io.StringIO()):
            first=trial.main(['--limit','3','--out',str(root/'first')])
            second=trial.main(['--limit','3','--out',str(root/'second')])
        self.assertEqual((first,second),(0,0));self.assertEqual(capture.call_count,2)
        self.assertFalse((root/'confirmed_owners.json').exists())
        self.assertEqual(len(list(root.rglob('*.xlsx'))),2)

    def test_cancelled_after_default_capture_does_not_start_any_query_worker(self):
        import contextlib,io
        from desktop_demo import owner_discovery,settings
        root=Path(self.temp.name)
        self.auto_config=SimpleNamespace(wdms_exe=str(root/'Synthetic.exe'),username='synthetic-user',credential_target='synthetic-target')
        cancel=root/'cancel.request'
        def capture(args):
            self.automatic_capture(args);cancel.touch()
            return 0
        with mock.patch.object(settings,'load_settings',return_value=self.auto_config), \
             mock.patch.object(owner_discovery,'main',side_effect=capture), \
             mock.patch.object(trial.probe,'supervise') as worker, \
             mock.patch.object(trial,'show_notice'),contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','3','--out',str(root/'out'),'--cancel-file',str(cancel)])
        self.assertEqual(code,2);worker.assert_not_called()
        self.assertEqual(list(root.rglob('*.xlsx')),[])

    def test_worker_rejects_changed_snapshot_before_native_api_or_query(self):
        import hashlib
        digest=hashlib.sha256(self.owners.read_bytes()).hexdigest()
        self.owners.write_text(json.dumps({'owners':list(reversed(self.owner_names))}), encoding='utf-8')
        with mock.patch.object(trial,'TrialNativeWindows') as api,mock.patch.object(trial.probe,'test_current_owner_query') as query:
            code=trial.child(SimpleNamespace(log=str(self.log),owners=self.owners,owners_sha256=digest))
        self.assertEqual(code,2);api.assert_not_called();query.assert_not_called()
        self.assertEqual(trial.probe.read_events(self.log)[-1]['error_code'],'OWNER_PLAN_CHANGED')

    def test_worker_rejects_different_pid_or_main_handle_before_selection(self):
        root=Path(self.temp.name)
        config=SimpleNamespace(wdms_exe=str(root/'Synthetic.exe'),username='synthetic-user',credential_target='synthetic-target')
        snapshot=root/'snapshot.json';owner_catalog.write_snapshot(snapshot,self.owner_names,config,{'pid':7,'hwnd':100})
        for pid,hwnd in ((8,100),(7,101)):
            with self.subTest(pid=pid,hwnd=hwnd):
                api=mock.Mock();api.find_main.return_value=[{'pid':pid,'hwnd':hwnd,'minimized':False}]
                with mock.patch.object(trial.probe,'inspect_environment'),mock.patch.object(trial.probe,'require_usable_environment'), \
                     mock.patch.object(trial,'TrialNativeWindows',return_value=api), \
                     mock.patch.object(trial,'select_owner') as select,mock.patch.object(trial.probe,'test_current_owner_query') as query:
                    code=trial.child(SimpleNamespace(log=str(self.log),owners=snapshot,title='synthetic-title'))
                self.assertEqual(code,2);select.assert_not_called();query.assert_not_called();api.read_owner_list.assert_not_called()

    def test_modal_disabled_session_after_capture_cannot_select_or_query(self):
        root=Path(self.temp.name)
        config=SimpleNamespace(wdms_exe=str(root/'Synthetic.exe'),username='synthetic-user',credential_target='synthetic-target')
        snapshot=root/'snapshot.json';owner_catalog.write_snapshot(snapshot,self.owner_names,config,{'pid':7,'hwnd':100})
        api=mock.Mock();api.find_main.return_value=[{'pid':7,'hwnd':100,'minimized':False}]
        api.user32.IsWindowEnabled.return_value=False
        with mock.patch.object(trial.probe,'inspect_environment'),mock.patch.object(trial.probe,'require_usable_environment'), \
             mock.patch.object(trial,'TrialNativeWindows',return_value=api), \
             mock.patch.object(trial,'select_owner') as select,mock.patch.object(trial.probe,'test_current_owner_query') as query:
            code=trial.child(SimpleNamespace(log=str(self.log),owners=snapshot,title='synthetic-title'))
        self.assertEqual(code,2);select.assert_not_called();query.assert_not_called()

    def api(self, name='目标货主'):
        self.notifications=[]
        def send(hwnd,msg,wparam,lparam,flags,timeout,result):
            self.notifications.append((hwnd,msg,wparam,lparam))
            result._obj.value=0
            return True
        return SimpleNamespace(metadata=lambda hwnd: {'pid':10,'class_name':'COMBO','visible':True},
             health=lambda *a:None, read_label_text=lambda *a:name,
             user32=SimpleNamespace(GetParent=lambda h:42, GetDlgCtrlID=lambda h:67284, SendMessageTimeoutW=send))

    def test_selection_notifies_once_and_verifies_index_and_name(self):
        api=self.api()
        with mock.patch.object(trial,'combo_send',side_effect=[0,1,1]):
            trial.select_owner(api,{'hwnd':10,'pid':10},{'hwnd':2,'class_name':'COMBO'},1,'目标货主',self.log)
        self.assertEqual(self.notifications,[(42,0x111,(1<<16)|(67284 & 0xffff),2)])

    def test_already_selected_owner_gets_no_notification(self):
        api=self.api()
        with mock.patch.object(trial,'combo_send',side_effect=[1,1]):
            trial.select_owner(api,{'hwnd':10,'pid':10},{'hwnd':2,'class_name':'COMBO'},1,'目标货主',self.log)
        self.assertEqual(self.notifications,[])

    def test_wrong_selection_stops_without_notification(self):
        api=self.api()
        with mock.patch.object(trial,'combo_send',side_effect=[0,-1]):
            with self.assertRaises(WDMSError):
                trial.select_owner(api,{'hwnd':10,'pid':10},{'hwnd':2,'class_name':'COMBO'},1,'目标货主',self.log)
        self.assertEqual(self.notifications,[])

    def test_wrong_owner_name_cannot_complete_selection(self):
        api=self.api(name='错误货主')
        with mock.patch.object(trial,'combo_send',return_value=1), mock.patch.object(trial.time,'monotonic',side_effect=[0,0,4]),mock.patch.object(trial.time,'sleep'):
            with self.assertRaises(WDMSError) as caught:
                trial.select_owner(api,{'hwnd':10,'pid':10},{'hwnd':2,'class_name':'COMBO'},1,'目标货主',self.log)
        self.assertEqual(caught.exception.code,'OWNER_SELECTION_MISMATCH')

    def test_excel_preserves_fraction_and_marks_manual_review(self):
        from openpyxl import load_workbook
        path=Path(self.temp.name)/'test.xlsx'
        trial.write_trial_excel([{'owner':'目标货主','whole':'261534','loose':'1304742.54','total':'32672335.54','time':'now'}],path)
        book=load_workbook(path,data_only=True)
        self.assertEqual(book.active.cell(2,3).value,1304742.54)
        self.assertEqual(book.active.max_column,3)
        self.assertEqual([c.value for c in book.active[1]],['货主','整件数','散件数'])
        self.assertEqual(book.active.title,trial_excel_reporter.inventory_report_title(date.today()))
        book.close()

    def test_query_wait_cursor_waits_for_idle_instead_of_aborting(self):
        states=iter([True,True,False])
        api=SimpleNamespace(guard_cursor=lambda:None, cursor_state=lambda:{'standard_wait_cursor':next(states),'cursor_visible':True})
        clock=[0.0]
        with mock.patch.object(trial.time,'monotonic',side_effect=lambda:clock[0]), \
             mock.patch.object(trial.time,'sleep',side_effect=lambda seconds:clock.__setitem__(0,clock[0]+seconds)):
            waited=trial.wait_until_idle(api,self.log,'目标货主')
        self.assertAlmostEqual(waited,0.4,places=6)
        events=[e['event'] for e in trial.probe.read_events(self.log)]
        self.assertIn('query_idle_wait',events)
        self.assertIn('query_idle_recovered',events)
        self.assertNotIn('query_idle_wait_timeout',events)

    def test_query_wait_cursor_that_never_clears_still_stops(self):
        api=SimpleNamespace(guard_cursor=lambda:None, cursor_state=lambda:{'standard_wait_cursor':True,'cursor_visible':True})
        clock=[0.0]
        with mock.patch.object(trial.time,'monotonic',side_effect=lambda:clock[0]), \
             mock.patch.object(trial.time,'sleep',side_effect=lambda seconds:clock.__setitem__(0,clock[0]+seconds)):
            with self.assertRaises(WDMSError) as caught:
                trial.wait_until_idle(api,self.log,'目标货主',limit_seconds=1.0)
        self.assertEqual(caught.exception.code,'QUERY_ALREADY_RUNNING')
        events=[e['event'] for e in trial.probe.read_events(self.log)]
        self.assertIn('query_idle_wait_timeout',events)

    def test_first_owner_failure_stops_batch_without_excel(self):
        import contextlib
        import io
        out=Path(self.temp.name)/'batch'
        with mock.patch.object(trial.probe,'supervise',return_value=2) as supervise, mock.patch.object(trial,'write_trial_excel') as writer, contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','3','--out',str(out),'--owners',str(self.owners)])
        self.assertEqual(code,2)
        self.assertEqual(supervise.call_count,1)
        writer.assert_not_called()
        self.notice_popup.assert_called_once()
        self.assertEqual(self.notice_popup.call_args.args[0], '库存采集已停止')
        self.assertIn('退出码为2', self.notice_popup.call_args.args[1])
        self.notice_factory.return_value.close.assert_called()


    def test_auto_start_requires_no_confirmation_dialog(self):
        import contextlib, io
        with mock.patch.object(trial,'show_notice',return_value=True) as notice, mock.patch.object(trial.probe,'supervise',return_value=2) as worker, contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','3','--out',self.temp.name,'--owners',str(self.owners)])
        self.assertEqual(code,2)
        self.assertEqual(worker.call_count,1)
        self.assertEqual([call.args[0] for call in notice.call_args_list],['库存采集已停止'])

    def test_unlogged_worker_exit_is_not_reported_as_a_recovered_read_timeout(self):
        import contextlib,io
        def worker(command,log,timeout,cancelled=None,stderr_path=None):
            self.assertEqual(stderr_path.name,'owner_01.stderr.txt')
            stderr_path.write_text('synthetic interpreter failure', encoding='utf-8')
            trial.probe.record(log,'query_settle_read_unavailable',error_code='NATIVE_TEXT_READ_FAILED',error='synthetic recovered read')
            return 7
        with mock.patch.object(trial.probe,'supervise',side_effect=worker), \
             mock.patch.object(trial,'show_notice') as notice,contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','3','--out',self.temp.name,'--owners',str(self.owners)])
        self.assertEqual(code,2)
        log=next(Path(self.temp.name).glob('native_trial_*/owner_01.jsonl'))
        error=trial.probe.read_events(log)[-1]
        self.assertEqual((error['event'],error['error_code'],error['exit_code']),('trial_worker_failed','TRIAL_WORKER_EXITED',7))
        self.assertIn('退出码为7',notice.call_args.args[1])
        self.assertEqual(list(Path(self.temp.name).rglob('*.xlsx')),[])

    def test_zero_exit_without_marker_is_explicitly_unconfirmed(self):
        import contextlib,io
        with mock.patch.object(trial.probe,'supervise',return_value=0), \
             mock.patch.object(trial,'show_notice'),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(trial.main(['--limit','1','--out',self.temp.name,'--owners',str(self.owners)]),2)
        log=next(Path(self.temp.name).glob('native_trial_*/owner_01.jsonl'))
        self.assertEqual(trial.probe.read_events(log)[-1]['error_code'],'TRIAL_WORKER_UNCONFIRMED')

    def test_terminal_error_blocks_worker_success_and_keeps_original_reason(self):
        import contextlib,io
        def worker(command,log,timeout,cancelled=None,stderr_path=None):
            publish_failure(command,'QUERY_OWNER_MISMATCH','synthetic owner mismatch')
            trial.probe.record(log,'probe_error',error_code='QUERY_OWNER_MISMATCH',error='synthetic owner mismatch')
            return 0
        with mock.patch.object(trial.probe,'supervise',side_effect=worker), \
             mock.patch.object(trial,'show_notice') as notice,contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(trial.main(['--limit','1','--out',self.temp.name,'--owners',str(self.owners)]),2)
        self.assertIn('synthetic owner mismatch',notice.call_args.args[1])
        self.assertFalse(list(Path(self.temp.name).rglob('*.xlsx')))

    def test_completion_notice_is_after_report_saved(self):
        import contextlib, io
        notices=[]
        def notice(title,message,**kwargs):
            if title=='库存采集完成':
                self.assertEqual(len(list(Path(self.temp.name).rglob('*.xlsx'))),1)
                self.assertEqual(len(list(Path(self.temp.name).rglob('trial_results.json'))),1)
            notices.append(title)
            return True
        with mock.patch.object(trial,'show_notice',side_effect=notice), mock.patch.object(trial.probe,'supervise',side_effect=self.successful_worker), contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','1','--out',self.temp.name,'--owners',str(self.owners)])
        self.assertEqual(code,0)
        self.assertEqual(notices,['库存采集完成'])

    def test_failed_report_never_shows_completion_notice(self):
        import contextlib, io
        with mock.patch.object(trial,'show_notice',return_value=True) as notice, mock.patch.object(trial.probe,'supervise',side_effect=self.successful_worker), mock.patch.object(trial,'write_trial_excel',side_effect=OSError('test')), contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','1','--out',self.temp.name,'--owners',str(self.owners)])
        self.assertEqual(code,2)
        self.assertEqual([call.args[0] for call in notice.call_args_list],['库存采集未完成'])

    def test_transient_health_timeout_recovers_without_query(self):
        api=object.__new__(trial.TrialNativeWindows)
        with mock.patch.object(trial.probe.NativeWindows,'health',side_effect=[WDMSError('WDMS_HEALTH_CHECK_FAILED','timeout'),None]) as health, mock.patch.object(trial.time,'sleep'):
            api.health(10,self.log,'initial')
        self.assertEqual(health.call_count,2)
        self.assertTrue(any(e['event']=='health_recovered' for e in trial.probe.read_events(self.log)))

    def test_health_wait_has_a_finite_deadline(self):
        api=object.__new__(trial.TrialNativeWindows)
        with mock.patch.object(trial.probe.NativeWindows,'health',side_effect=WDMSError('WDMS_HEALTH_CHECK_FAILED','timeout')) as health, mock.patch.object(trial.time,'monotonic',side_effect=[0,16]):
            with self.assertRaises(WDMSError) as caught:
                api.health(10,self.log,'initial')
        self.assertEqual(caught.exception.code,'WDMS_HEALTH_WAIT_TIMEOUT')
        self.assertEqual(health.call_count,1)

    def test_access_denied_is_not_retried(self):
        api=object.__new__(trial.TrialNativeWindows)
        with mock.patch.object(trial.probe.NativeWindows,'health',side_effect=WDMSError('WDMS_ACCESS_DENIED','denied')) as health:
            with self.assertRaises(WDMSError):
                api.health(10,self.log,'initial')
        self.assertEqual(health.call_count,1)

    def test_running_notice_closes_before_completion_popup(self):
        import contextlib, io
        banner=mock.Mock()
        def notice(title,message,**kwargs):
            banner.close.assert_called()
            return True
        with mock.patch.object(trial,'RunningNotice',return_value=banner), mock.patch.object(trial,'show_notice',side_effect=notice), mock.patch.object(trial.probe,'supervise',side_effect=self.successful_worker), contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','1','--out',self.temp.name,'--owners',str(self.owners)])
        self.assertEqual(code,0)
        banner.start.assert_called_once()
        banner.update.assert_called_once()

    def test_managed_collection_has_no_intermediate_popup_and_exports_manifest(self):
        import contextlib,io,json
        root=Path(self.temp.name)
        manifest=root/'collection_result.json'
        with mock.patch.object(trial,'show_notice') as popup,mock.patch.object(trial,'RunningNotice') as banner,mock.patch.object(trial.probe,'supervise',side_effect=self.successful_worker),contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','3','--out',str(root/'reports'),'--owners',str(self.owners),'--managed-notices','--workflow-progress',str(root/'progress.txt'),'--result-manifest',str(manifest)])
        self.assertEqual(code,0)
        popup.assert_not_called();banner.assert_not_called()
        result=json.loads(manifest.read_text(encoding='utf-8'))
        self.assertEqual(result['row_count'],3)
        self.assertTrue(Path(result['report']).is_file())
        self.assertIn('3/3',(root/'progress.txt').read_text(encoding='utf-8'))

    def test_cancel_before_first_worker_creates_no_manifest(self):
        import contextlib,io
        root=Path(self.temp.name)
        cancel=root/'cancel.request';cancel.touch()
        manifest=root/'collection_result.json'
        notice=mock.Mock()
        with mock.patch.object(trial,'RunningNotice',return_value=notice), \
             mock.patch.object(trial,'show_notice') as popup, \
             mock.patch.object(trial.probe,'supervise') as worker, contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','3','--out',str(root/'out'),'--owners',str(self.owners),
                             '--cancel-file',str(cancel),'--result-manifest',str(manifest)])
        self.assertEqual(code,2)
        worker.assert_not_called();notice.close.assert_called()
        popup.assert_called_once()
        self.assertEqual(popup.call_args.args[0],'库存采集已停止')
        self.assertFalse(manifest.exists())
        event=trial.probe.read_events(next((root/'out').glob('native_trial_*'))/'trial_error.jsonl')[-1]
        self.assertEqual(event['error_code'],'TASK_CANCELLED')

    def test_cancel_between_owners_stops_before_second_worker(self):
        import contextlib,io
        root=Path(self.temp.name)
        cancel=root/'cancel.request'
        manifest=root/'collection_result.json'
        calls=[]

        def worker(command,log,timeout,cancelled=None,stderr_path=None):
            calls.append(command)
            self.successful_worker(command,log,timeout,cancelled)
            cancel.touch()
            return 0

        with mock.patch.object(trial.probe,'supervise',side_effect=worker) as supervisor, \
             mock.patch.object(trial,'write_trial_excel') as writer, \
             mock.patch.object(trial,'show_notice') as popup,contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','3','--out',str(root/'out'),'--owners',str(self.owners),
                             '--cancel-file',str(cancel),'--result-manifest',str(manifest)])
        self.assertEqual(code,2)
        self.assertEqual(len(calls),1);self.assertEqual(supervisor.call_count,1)
        self.assertEqual(set(supervisor.call_args.kwargs),{'cancelled','stderr_path'})
        self.assertEqual(supervisor.call_args.kwargs['cancelled'](),True)
        writer.assert_not_called();self.assertFalse(manifest.exists())
        self.assertEqual(popup.call_args.args[0],'库存采集已停止')

    def test_cancel_during_excel_write_keeps_unconfirmed_workbook_without_manifest(self):
        import contextlib,io
        root=Path(self.temp.name)
        cancel=root/'cancel.request'
        manifest=root/'collection_result.json'
        notices=[]

        def write_and_cancel(rows,path,task_date=None):
            Path(path).write_bytes(b'workbook fixture')
            cancel.touch()

        with mock.patch.object(trial.probe,'supervise',side_effect=self.successful_worker), \
             mock.patch.object(trial,'write_trial_excel',side_effect=write_and_cancel), \
             mock.patch.object(trial,'show_notice',side_effect=lambda title,message,**kwargs:notices.append((title,message))), \
             contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','1','--out',str(root/'out'),'--owners',str(self.owners),
                             '--cancel-file',str(cancel),'--result-manifest',str(manifest)])
        self.assertEqual(code,2)
        self.assertTrue(list((root/'out').glob(f"native_trial_*/{trial_excel_reporter.inventory_report_title(date.today())}.xlsx")))
        self.assertFalse(manifest.exists())
        self.assertEqual(notices[-1][0],'库存采集已停止')
        self.assertIn('本轮结果未确认',notices[-1][1])

    def test_total_deadline_expiry_stops_without_manifest(self):
        import contextlib,io
        root=Path(self.temp.name)
        manifest=root/'collection_result.json'
        clock=[0.0]

        def worker(command,log,timeout,cancelled=None,stderr_path=None):
            self.successful_worker(command,log,timeout,cancelled)
            clock[0]=1.0
            return 0

        with mock.patch.object(trial.time,'monotonic',side_effect=lambda:clock[0]), \
             mock.patch.object(trial.probe,'supervise',side_effect=worker) as supervisor, \
             mock.patch.object(trial,'write_trial_excel') as writer, \
             mock.patch.object(trial,'show_notice') as popup,contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','1','--out',str(root/'out'),'--owners',str(self.owners),
                             '--execution-deadline','1','--result-manifest',str(manifest)])
        self.assertEqual(code,2)
        self.assertEqual(supervisor.call_args.args[2],1)
        writer.assert_not_called();self.assertFalse(manifest.exists())
        self.assertEqual(popup.call_args.args[0],'库存采集已停止')
        event=trial.probe.read_events(next((root/'out').glob('native_trial_*'))/'trial_error.jsonl')[-1]
        self.assertEqual(event['error_code'],'STEP_TIMEOUT')

    def test_stop_after_manifest_publication_removes_this_manifest(self):
        import contextlib,io
        root=Path(self.temp.name)
        manifest=root/'collection_result.json'
        clock=[0.0]
        original_replace=Path.replace

        def publish_and_expire(path,target):
            result=original_replace(path,target)
            if path.name=='collection_result.tmp':
                clock[0]=1.0
            return result

        with mock.patch.object(trial.time,'monotonic',side_effect=lambda:clock[0]), \
             mock.patch.object(trial.probe,'supervise',side_effect=self.successful_worker), \
             mock.patch.object(trial,'write_trial_excel',side_effect=lambda rows,path,task_date=None:Path(path).write_bytes(b'report')), \
             mock.patch.object(Path,'replace',new=publish_and_expire), \
             mock.patch.object(trial,'show_notice') as popup,contextlib.redirect_stdout(io.StringIO()):
            code=trial.main(['--limit','1','--out',str(root/'out'),'--owners',str(self.owners),
                             '--execution-deadline','1','--result-manifest',str(manifest)])
        self.assertEqual(code,2)
        self.assertFalse(manifest.exists())
        self.assertEqual(popup.call_args.args[0],'库存采集已停止')


class TrialExcelReporterTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def test_whole_count_is_strict_integer_and_fraction_is_preserved(self):
        with self.assertRaises(WDMSError) as caught:
            trial_excel_reporter.write_trial_excel(
                [{'owner':'A','whole':'1.5','loose':'2.75'}],self.root/'invalid.xlsx')
        self.assertEqual(caught.exception.code,'INVALID_COUNT')

        path=self.root/'valid.xlsx'
        trial_excel_reporter.write_trial_excel(
            [{'owner':'A','whole':'2','loose':'2.75'}],path)
        from openpyxl import load_workbook
        workbook=load_workbook(path,data_only=True)
        self.assertEqual(workbook.active['B2'].value,2)
        self.assertEqual(workbook.active['C2'].value,2.75)
        self.assertEqual(workbook.active.freeze_panes,'A2')
        self.assertEqual([workbook.active.column_dimensions[key].width for key in ('A','B','C')],[48,18,22])
        workbook.close()

    def test_empty_duplicate_and_blank_owner_inputs_are_rejected(self):
        cases=([],
               [{'owner':'A','whole':'1','loose':'0'},{'owner':'A','whole':'2','loose':'0'}],
               [{'owner':'  ','whole':'1','loose':'0'}],
               [None],
               [{'owner':'A','whole':'1'}])
        for rows in cases:
            with self.subTest(rows=rows):
                with self.assertRaises(WDMSError) as caught:
                    trial_excel_reporter.write_trial_excel(rows,self.root/'invalid.xlsx')
                self.assertEqual(caught.exception.code,'REPORT_INPUT_INVALID')

    def test_formula_like_owner_is_stored_as_a_string(self):
        path=self.root/'formula.xlsx'
        trial_excel_reporter.write_trial_excel(
            [{'owner':'=1+1','whole':'1','loose':'0'}],path)
        from openpyxl import load_workbook
        workbook=load_workbook(path,data_only=False)
        cell=workbook.active['A2']
        self.assertEqual(cell.value,'=1+1')
        self.assertEqual(cell.data_type,'s')
        workbook.close()

    def test_readback_mismatch_preserves_existing_target(self):
        path=self.root/'existing.xlsx'
        path.write_bytes(b'original target')
        with mock.patch.object(trial_excel_reporter,'_check_readback',side_effect=WDMSError('REPORT_READBACK_MISMATCH','mismatch')):
            with self.assertRaises(WDMSError) as caught:
                trial_excel_reporter.write_trial_excel(
                    [{'owner':'A','whole':'1','loose':'0'}],path)
        self.assertEqual(caught.exception.code,'REPORT_READBACK_MISMATCH')
        self.assertEqual(path.read_bytes(),b'original target')
        self.assertEqual(list(self.root.glob('*.pending.xlsx')),[])

    def test_save_failure_cleans_pending_file(self):
        path=self.root/'failed.xlsx'
        with mock.patch.object(trial_excel_reporter.Workbook,'save',side_effect=OSError('hidden save detail')):
            with self.assertRaises(WDMSError) as caught:
                trial_excel_reporter.write_trial_excel(
                    [{'owner':'A','whole':'1','loose':'0'}],path)
        self.assertEqual(caught.exception.code,'REPORT_PUBLISH_FAILED')
        self.assertEqual(list(self.root.glob('*.pending.xlsx')),[])
        self.assertFalse(path.exists())

if __name__=='__main__':unittest.main()
